"""流式读取检验XLSX并写出受控的L1 Parquet分片。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import tempfile
import time
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, Iterator, Sequence

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq

from .lab_l1_schema import (
    DEFAULT_BATCH_SIZE,
    IDENTIFIER_FIELDS,
    INGESTION_FIELDS,
    INGESTION_VERSION,
    PILOT_ROW_LIMIT,
    PILOT_SHEETS,
    RAW_FIELDS,
    TIMESTAMP_FIELDS,
)


_SCIENTIFIC_INTEGER_RE = re.compile(r"^[+-]?\d+(?:\.0+)?[eE][+-]?\d+$")


@dataclass(frozen=True)
class ParquetShardSummary:
    path: str
    source_row_start: int
    source_row_end: int
    row_count: int
    byte_size: int
    sha256: str


@dataclass(frozen=True)
class SheetIngestionSummary:
    source_sheet: str
    requested_row_limit: int
    emitted_row_count: int
    source_row_start: int | None
    source_row_end: int | None
    reached_row_limit: bool
    null_counts: dict[str, int]
    duplicate_row_hash_count: int
    identifier_source_type_counts: dict[str, dict[str, int]]
    identifier_scientific_notation_output_count: int
    shards: tuple[ParquetShardSummary, ...]
    elapsed_seconds: float
    peak_memory_bytes: int


def _memory_rss_bytes() -> int:
    try:
        import psutil  # type: ignore

        return int(psutil.Process(os.getpid()).memory_info().rss)
    except Exception:
        return 0


def _jsonable_cell(value: Any) -> dict[str, Any]:
    if value is None:
        return {"type": "null", "value": None}
    if isinstance(value, datetime):
        return {"type": "datetime", "value": value.isoformat()}
    if isinstance(value, date):
        return {"type": "date", "value": value.isoformat()}
    if isinstance(value, bool):
        return {"type": "bool", "value": value}
    if isinstance(value, int):
        return {"type": "int", "value": value}
    if isinstance(value, float):
        if math.isnan(value):
            serial_value: Any = "NaN"
        elif math.isinf(value):
            serial_value = "Infinity" if value > 0 else "-Infinity"
        else:
            serial_value = repr(value)
        return {"type": "float", "value": serial_value}
    return {"type": type(value).__name__, "value": str(value)}


def row_hash(raw_values: Sequence[Any]) -> str:
    payload = json.dumps(
        [_jsonable_cell(value) for value in raw_values],
        ensure_ascii=False,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def _identifier_to_string(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if math.isnan(value):
            return "NaN"
        if math.isinf(value):
            return "Infinity" if value > 0 else "-Infinity"
        # Fixed-point formatting prevents scientific notation in the output.
        text = format(value, "f")
        if "." in text:
            text = text.rstrip("0").rstrip(".")
        return text
    return str(value)


def _text_value(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        return value
    raise TypeError(f"non-text value in a text laboratory field: {type(value).__name__}")


def _timestamp_value(value: Any) -> datetime | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    raise TypeError(f"non-datetime value in a laboratory time field: {type(value).__name__}")


def _normalize_row(raw_values: Sequence[Any]) -> tuple[list[Any], str]:
    if len(raw_values) != len(RAW_FIELDS):
        raise ValueError(f"expected {len(RAW_FIELDS)} source fields, got {len(raw_values)}")
    normalized: list[Any] = []
    for field, value in zip(RAW_FIELDS, raw_values):
        if field in IDENTIFIER_FIELDS:
            normalized.append(_identifier_to_string(value))
        elif field in TIMESTAMP_FIELDS:
            normalized.append(_timestamp_value(value))
        else:
            normalized.append(_text_value(value))
    return normalized, row_hash(raw_values)


def _parquet_schema() -> pa.Schema:
    fields: list[pa.Field] = []
    for field in RAW_FIELDS:
        if field in IDENTIFIER_FIELDS:
            fields.append(pa.field(field, pa.string()))
        elif field in TIMESTAMP_FIELDS:
            fields.append(pa.field(field, pa.timestamp("us")))
        else:
            fields.append(pa.field(field, pa.large_string()))
    fields.extend(
        [
            pa.field("source_workbook", pa.string()),
            pa.field("source_sheet", pa.string()),
            pa.field("source_row", pa.int64()),
            pa.field("source_record_id", pa.string()),
            pa.field("row_hash", pa.string()),
            pa.field("ingestion_version", pa.string()),
        ]
    )
    return pa.schema(fields)


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _write_parquet_shard(
    output_path: Path,
    source_workbook: str,
    source_sheet: str,
    rows: Sequence[tuple[int, Sequence[Any], Sequence[Any], str]],
) -> ParquetShardSummary:
    normalized_rows: list[list[Any]] = []
    for source_row, raw_values, normalized_values, digest in rows:
        normalized_rows.append(
            list(normalized_values)
            + [
                source_workbook,
                source_sheet,
                source_row,
                f"{source_sheet}:{source_row}",
                digest,
                INGESTION_VERSION,
            ]
        )
    columns = list(zip(*normalized_rows)) if normalized_rows else []
    table = pa.Table.from_arrays(
        [pa.array(values, type=field.type) for values, field in zip(columns, _parquet_schema())],
        schema=_parquet_schema(),
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(prefix=output_path.name + ".", suffix=".tmp", dir=str(output_path.parent))
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        pq.write_table(table, temp_path, compression="zstd")
        temp_path.replace(output_path)
    except Exception:
        temp_path.unlink(missing_ok=True)
        raise
    return ParquetShardSummary(
        path=str(output_path),
        source_row_start=rows[0][0],
        source_row_end=rows[-1][0],
        row_count=len(rows),
        byte_size=output_path.stat().st_size,
        sha256=_sha256_file(output_path),
    )


def _validate_header(ws: Any, source_sheet: str) -> None:
    header = tuple(next(ws.iter_rows(min_row=1, max_row=1, max_col=len(RAW_FIELDS), values_only=True)))
    if header != RAW_FIELDS:
        raise ValueError(f"unexpected header in {source_sheet}: {header!r}")


def _iter_source_rows(ws: Any, row_limit: int) -> Iterator[tuple[int, tuple[Any, ...]]]:
    for source_row, values in enumerate(
        ws.iter_rows(min_row=2, max_row=row_limit + 1, max_col=len(RAW_FIELDS), values_only=True),
        start=2,
    ):
        yield source_row, tuple(values)


def _ingest_sheet(
    ws: Any,
    *,
    source_workbook: str,
    output_dir: Path,
    row_limit: int,
    batch_size: int,
    peak_memory_start: int,
) -> SheetIngestionSummary:
    source_sheet = ws.title
    _validate_header(ws, source_sheet)
    started = time.perf_counter()
    peak_memory = max(peak_memory_start, _memory_rss_bytes())
    null_counts = Counter()
    identifier_types: dict[str, Counter[str]] = {field: Counter() for field in IDENTIFIER_FIELDS}
    identifier_scientific_output_count = 0
    row_hash_counts: Counter[str] = Counter()
    shards: list[ParquetShardSummary] = []
    batch: list[tuple[int, Sequence[Any], Sequence[Any], str]] = []
    emitted = 0
    first_source_row: int | None = None
    last_source_row: int | None = None

    for source_row, raw_values in _iter_source_rows(ws, row_limit):
        normalized_values, digest = _normalize_row(raw_values)
        for field, value in zip(RAW_FIELDS, raw_values):
            if value is None:
                null_counts[field] += 1
        for field, value in zip(RAW_FIELDS, raw_values):
            if field in IDENTIFIER_FIELDS:
                identifier_types[field][type(value).__name__] += 1
        for field, raw_value, value in zip(RAW_FIELDS, raw_values, normalized_values):
            if (
                field in IDENTIFIER_FIELDS
                and isinstance(raw_value, float)
                and isinstance(value, str)
                and _SCIENTIFIC_INTEGER_RE.fullmatch(value)
            ):
                identifier_scientific_output_count += 1
        row_hash_counts[digest] += 1
        batch.append((source_row, raw_values, normalized_values, digest))
        emitted += 1
        first_source_row = source_row if first_source_row is None else first_source_row
        last_source_row = source_row
        if len(batch) >= batch_size:
            shard_path = output_dir / f"part-{len(shards) + 1:05d}.parquet"
            shards.append(_write_parquet_shard(shard_path, source_workbook, source_sheet, batch))
            batch.clear()
            peak_memory = max(peak_memory, _memory_rss_bytes())
    if batch:
        shard_path = output_dir / f"part-{len(shards) + 1:05d}.parquet"
        shards.append(_write_parquet_shard(shard_path, source_workbook, source_sheet, batch))
        batch.clear()
        peak_memory = max(peak_memory, _memory_rss_bytes())

    peak_memory = max(peak_memory, _memory_rss_bytes())
    return SheetIngestionSummary(
        source_sheet=source_sheet,
        requested_row_limit=row_limit,
        emitted_row_count=emitted,
        source_row_start=first_source_row,
        source_row_end=last_source_row,
        reached_row_limit=emitted == row_limit,
        null_counts=dict(sorted(null_counts.items())),
        duplicate_row_hash_count=sum(count - 1 for count in row_hash_counts.values() if count > 1),
        identifier_source_type_counts={field: dict(sorted(counts.items())) for field, counts in sorted(identifier_types.items())},
        identifier_scientific_notation_output_count=identifier_scientific_output_count,
        shards=tuple(shards),
        elapsed_seconds=time.perf_counter() - started,
        peak_memory_bytes=peak_memory,
    )


def ingest_workbook(
    workbook_path: str | Path,
    output_root: str | Path,
    *,
    sheet_names: Sequence[str] = PILOT_SHEETS,
    row_limit: int = PILOT_ROW_LIMIT,
    batch_size: int = DEFAULT_BATCH_SIZE,
    report_path: str | Path | None = None,
) -> dict[str, Any]:
    if row_limit <= 0 or batch_size <= 0 or batch_size > row_limit:
        raise ValueError("batch_size and row_limit must be positive, with batch_size <= row_limit")
    requested_sheets = tuple(sheet_names)
    if not requested_sheets or any(name not in PILOT_SHEETS for name in requested_sheets):
        raise ValueError(f"only the pilot sheets are allowed: {PILOT_SHEETS!r}")
    if len(set(requested_sheets)) != len(requested_sheets):
        raise ValueError("duplicate sheet name in request")

    workbook = Path(workbook_path).resolve()
    output_root_path = Path(output_root).resolve()
    report_target = Path(report_path).resolve() if report_path else output_root_path.parent / "lab_ingestion_report.json"
    if output_root_path.exists() and any(output_root_path.iterdir()):
        raise FileExistsError(f"refusing to use non-empty output directory: {output_root_path}")
    output_root_path.mkdir(parents=True, exist_ok=True)
    started = time.perf_counter()
    peak_memory = _memory_rss_bytes()
    workbook_size = workbook.stat().st_size
    summaries: list[SheetIngestionSummary] = []

    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    try:
        unknown = [name for name in requested_sheets if name not in wb.sheetnames]
        if unknown:
            raise ValueError(f"requested sheet not found: {unknown!r}")
        untouched_sheets = [name for name in wb.sheetnames if name not in requested_sheets]
        for sheet_name in requested_sheets:
            summary = _ingest_sheet(
                wb[sheet_name],
                source_workbook=str(workbook),
                output_dir=output_root_path / sheet_name,
                row_limit=row_limit,
                batch_size=batch_size,
                peak_memory_start=peak_memory,
            )
            summaries.append(summary)
            peak_memory = max(peak_memory, summary.peak_memory_bytes)
    finally:
        wb.close()

    report: dict[str, Any] = {
        "report_version": "lab_ingestion_report_v1",
        "ingestion_version": INGESTION_VERSION,
        "source_workbook": str(workbook),
        "source_workbook_size": workbook_size,
        "read_mode": {"openpyxl_read_only": True, "data_only": True, "pandas_used": False, "cross_sheet_concat": False},
        "requested_sheets": list(requested_sheets),
        "untouched_sheets": untouched_sheets,
        "row_limit_excludes_header": True,
        "row_limit": row_limit,
        "batch_size": batch_size,
        "total_emitted_row_count": sum(summary.emitted_row_count for summary in summaries),
        "peak_memory_bytes": peak_memory,
        "peak_memory_gib": peak_memory / (1024**3),
        "elapsed_seconds": time.perf_counter() - started,
        "identifier_output_types": {field: "string" for field in sorted(IDENTIFIER_FIELDS)},
        "identifier_scientific_notation_output_count": sum(
            summary.identifier_scientific_notation_output_count for summary in summaries
        ),
        "sheets": [
            {
                **asdict(summary),
                "shards": [asdict(shard) for shard in summary.shards],
            }
            for summary in summaries
        ],
    }
    report_target.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(prefix=report_target.name + ".", suffix=".tmp", dir=str(report_target.parent))
    os.close(fd)
    temp_report = Path(temp_name)
    try:
        temp_report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        temp_report.replace(report_target)
    except Exception:
        temp_report.unlink(missing_ok=True)
        raise
    report["report_path"] = str(report_target)
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--output-root", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    parser.add_argument("--row-limit", type=int, default=PILOT_ROW_LIMIT)
    parser.add_argument("--batch-size", type=int, default=DEFAULT_BATCH_SIZE)
    args = parser.parse_args()
    report = ingest_workbook(
        args.workbook,
        args.output_root,
        row_limit=args.row_limit,
        batch_size=args.batch_size,
        report_path=args.report,
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
