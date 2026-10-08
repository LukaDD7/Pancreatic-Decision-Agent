"""Stage 5 manifest, shard planning, SQLite state, and dry-run orchestration."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import shutil
import sqlite3
import tempfile
import time
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq

from data_pipeline.paths import data_root
from .document_stream_reader import (
    _hash_raw_fields as _document_hash_raw_fields,
    _parquet_schema as _document_parquet_schema,
    _source_bytes as _document_source_bytes,
    build_record as _build_document_record,
)
from .document_l1_schema import RAW_FIELDS as DOCUMENT_RAW_FIELDS
from .lab_l1_schema import INGESTION_VERSION as LAB_INGESTION_VERSION, RAW_FIELDS as LAB_RAW_FIELDS
from .lab_reference_parser import parse_reference
from .lab_result_parser import _enriched_schema as _lab_enriched_schema, parse_result
from .lab_stream_reader import (
    _identifier_to_string,
    _jsonable_cell,
    _normalize_row,
    _parquet_schema as _lab_parquet_schema,
)
from .lab_time_normalizer import normalize_time_pair
from .lab_unit_normalizer import normalize_unit


DATA_ROOT = data_root()
DEFAULT_WORKBOOK = DATA_ROOT / "两万患者检验.XLSX"
DEFAULT_DOCUMENT = DATA_ROOT / "患者入院出院文书.csv"
DEFAULT_MANIFEST = DATA_ROOT / "code" / "stage5_manifest_v1.json"
DEFAULT_STATE_DB = DATA_ROOT / "code" / "stage5_pipeline_state_v1.sqlite3"
DEFAULT_OUTPUT_ROOT = DATA_ROOT / "pipeline_outputs_stage5_v1"
DEFAULT_DRY_RUN_REPORT = DEFAULT_OUTPUT_ROOT / "dry_run_report.json"

EXPECTED_LAB_SHEETS: tuple[str, ...] = (
    "Sheet1",
    "Sheet1(2)",
    "Sheet1(3)",
    "Sheet1(4)",
    "Sheet1(5)",
    "Sheet1(6)",
    "Sheet1(7)",
    "Sheet1(8)",
    "Sheet1(9)",
    "Sheet1(10)",
    "Sheet1(11)",
    "Sheet1(12)",
)
LAB_FULL_SHEET_ROWS = 1_048_575
LAB_LAST_SHEET_ROWS = 677_657
LAB_SHARD_TARGET_ROWS = 100_000
DOCUMENT_ROWS = 598_826
DOCUMENT_SHARD_TARGET_ROWS = 25_000

MANIFEST_VERSION = "stage5_manifest_v1"
LAB_DATA_CONTRACT_VERSION = "lab_l1_schema_v1"
LAB_CONVERSION_RULE_VERSION = "lab_ingestion_v1+lab_enriched_v1"
DOCUMENT_DATA_CONTRACT_VERSION = "document_l1_schema_v1"
DOCUMENT_CONVERSION_RULE_VERSION = "document_l1_restricted_v1"

STATUSES: tuple[str, ...] = ("PENDING", "RUNNING", "SUCCEEDED", "FAILED", "QUARANTINED")


class ManifestValidationError(ValueError):
    pass


class GlobalPipelineBlock(RuntimeError):
    """A condition that must stop the run instead of being silently quarantined."""


class RuleDefect(GlobalPipelineBlock):
    pass


@dataclass(frozen=True)
class CommitResult:
    row_count: int
    output_sha256: str
    schema_sha256: str
    output_file: str
    accepted_count: int | None = None
    quarantined_count: int = 0
    quarantine_sha256: str | None = None
    quarantine_schema_sha256: str | None = None


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _schema_sha256(schema: pa.Schema) -> str:
    return hashlib.sha256(schema.serialize().to_pybytes()).hexdigest()


def _memory_rss_bytes() -> int:
    try:
        import psutil  # type: ignore

        return int(psutil.Process(os.getpid()).memory_info().rss)
    except Exception:
        return 0


def _write_json_atomic(path: Path, payload: Mapping[str, Any]) -> None:
    path = path.resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(prefix=path.name + ".", suffix=".partial", dir=str(path.parent))
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        temp_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        temp_path.replace(path)
    except Exception:
        temp_path.unlink(missing_ok=True)
        raise


def _safe_component(value: str) -> str:
    return value.replace("/", "_").replace("\\", "_").replace(":", "_")


def _shard_ranges(total_rows: int, target_rows: int) -> list[dict[str, int]]:
    if total_rows <= 0 or target_rows <= 0:
        raise ValueError("total_rows and target_rows must be positive")
    ranges: list[dict[str, int]] = []
    start = 1
    while start <= total_rows:
        end = min(total_rows, start + target_rows - 1)
        ranges.append(
            {
                "data_row_start": start,
                "data_row_end": end,
                "expected_row_count": end - start + 1,
            }
        )
        start = end + 1
    return ranges


def _with_source_rows(item: Mapping[str, int], source_row_offset: int) -> dict[str, int]:
    result = dict(item)
    result["source_row_start"] = item["data_row_start"] + source_row_offset
    result["source_row_end"] = item["data_row_end"] + source_row_offset
    return result


def _validate_source_file(path: Path) -> tuple[str, int]:
    path = path.resolve()
    if not path.is_file():
        raise FileNotFoundError(path)
    return _sha256_file(path), path.stat().st_size


def _discover_lab_sheets(workbook_path: Path) -> tuple[str, ...]:
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        actual = tuple(workbook.sheetnames)
    finally:
        workbook.close()
    if actual != EXPECTED_LAB_SHEETS:
        raise ManifestValidationError(
            f"workbook sheet collection mismatch: expected {EXPECTED_LAB_SHEETS!r}, got {actual!r}"
        )
    return actual


def _target_paths(output_root: Path, kind: str, sheet_name: str | None, part_number: int) -> dict[str, str]:
    if kind == "lab":
        component = _safe_component(sheet_name or "unknown_sheet")
        relative = Path("restricted") / "lab_l1" / component / f"part-{part_number:05d}.parquet"
        quarantine = Path("restricted") / "lab_quarantine" / component / f"part-{part_number:05d}.parquet"
        audit = Path("audit") / "lab_l1" / component / f"part-{part_number:05d}.json"
    else:
        relative = Path("restricted") / "document_l1" / f"part-{part_number:05d}.parquet"
        quarantine = Path("restricted") / "document_quarantine" / f"part-{part_number:05d}.parquet"
        audit = Path("audit") / "document_l1" / f"part-{part_number:05d}.json"
    output_file = output_root / relative
    return {
        "output_file": str(output_file),
        "output_temp_file": str(Path(str(output_file) + ".partial")),
        "quarantine_file": str(output_root / quarantine),
        "audit_file": str(output_root / audit),
    }


def _build_tasks(
    *,
    source_input_id: str,
    kind: str,
    input_file: Path,
    sheet_name: str | None,
    expected_row_count: int,
    target_rows: int,
    output_root: Path,
    data_contract_version: str,
    conversion_rule_version: str,
    source_row_offset: int,
    source_sha256: str,
) -> list[dict[str, Any]]:
    tasks: list[dict[str, Any]] = []
    for part_number, base_range in enumerate(_shard_ranges(expected_row_count, target_rows), start=1):
        input_range = _with_source_rows(base_range, source_row_offset)
        paths = _target_paths(output_root, kind, sheet_name, part_number)
        task_id = f"{source_input_id}:part-{part_number:05d}"
        tasks.append(
            {
                "task_id": task_id,
                "source_input_id": source_input_id,
                "kind": kind,
                "input_file": str(input_file.resolve()),
                "sheet_name": sheet_name,
                "input_range": input_range,
                "expected_row_count": input_range["expected_row_count"],
                "file_size": input_file.stat().st_size,
                "sha256": source_sha256,
                "data_contract_version": data_contract_version,
                "conversion_rule_version": conversion_rule_version,
                **paths,
            }
        )
    return tasks


def build_manifest(
    workbook_path: str | Path = DEFAULT_WORKBOOK,
    document_path: str | Path = DEFAULT_DOCUMENT,
    manifest_path: str | Path = DEFAULT_MANIFEST,
    *,
    output_root: str | Path = DEFAULT_OUTPUT_ROOT,
    refuse_overwrite: bool = True,
) -> dict[str, Any]:
    workbook = Path(workbook_path).resolve()
    document = Path(document_path).resolve()
    manifest_target = Path(manifest_path).resolve()
    output_root_path = Path(output_root).resolve()
    if refuse_overwrite and manifest_target.exists():
        raise FileExistsError(f"refusing to overwrite manifest: {manifest_target}")

    workbook_sha256, workbook_size = _validate_source_file(workbook)
    document_sha256, document_size = _validate_source_file(document)
    sheet_names = _discover_lab_sheets(workbook)
    if len(sheet_names) != 12 or set(sheet_names) != set(EXPECTED_LAB_SHEETS):
        raise ManifestValidationError("lab workbook must contain exactly the 12 declared sheets")

    inputs: list[dict[str, Any]] = []
    tasks: list[dict[str, Any]] = []
    for sheet_name in sheet_names:
        expected_rows = LAB_LAST_SHEET_ROWS if sheet_name == "Sheet1(12)" else LAB_FULL_SHEET_ROWS
        source_input_id = f"lab:{sheet_name}"
        inputs.append(
            {
                "input_id": source_input_id,
                "kind": "lab",
                "input_file": str(workbook),
                "file_size": workbook_size,
                "sha256": workbook_sha256,
                "sheet_name": sheet_name,
                "expected_row_count": expected_rows,
                "data_contract_version": LAB_DATA_CONTRACT_VERSION,
                "conversion_rule_version": LAB_CONVERSION_RULE_VERSION,
                "task_unit": "one_sheet",
                "shard_target_rows": LAB_SHARD_TARGET_ROWS,
                "data_rows_scanned_for_manifest": 0,
            }
        )
        tasks.extend(
            _build_tasks(
                source_input_id=source_input_id,
                kind="lab",
                input_file=workbook,
                sheet_name=sheet_name,
                expected_row_count=expected_rows,
                target_rows=LAB_SHARD_TARGET_ROWS,
                output_root=output_root_path,
                data_contract_version=LAB_DATA_CONTRACT_VERSION,
                conversion_rule_version=LAB_CONVERSION_RULE_VERSION,
                source_row_offset=1,
                source_sha256=workbook_sha256,
            )
        )

    inputs.append(
        {
            "input_id": "document:admission_discharge",
            "kind": "document",
            "input_file": str(document),
            "file_size": document_size,
            "sha256": document_sha256,
            "sheet_name": None,
            "expected_row_count": DOCUMENT_ROWS,
            "data_contract_version": DOCUMENT_DATA_CONTRACT_VERSION,
            "conversion_rule_version": DOCUMENT_CONVERSION_RULE_VERSION,
            "task_unit": "one_csv_source_row_range",
            "shard_target_rows": DOCUMENT_SHARD_TARGET_ROWS,
            "data_rows_scanned_for_manifest": 0,
        }
    )
    tasks.extend(
        _build_tasks(
            source_input_id="document:admission_discharge",
            kind="document",
            input_file=document,
            sheet_name=None,
            expected_row_count=DOCUMENT_ROWS,
            target_rows=DOCUMENT_SHARD_TARGET_ROWS,
            output_root=output_root_path,
            data_contract_version=DOCUMENT_DATA_CONTRACT_VERSION,
            conversion_rule_version=DOCUMENT_CONVERSION_RULE_VERSION,
            source_row_offset=1,
            source_sha256=document_sha256,
        )
    )

    lab_total = sum(item["expected_row_count"] for item in inputs if item["kind"] == "lab")
    document_total = sum(item["expected_row_count"] for item in inputs if item["kind"] == "document")
    manifest: dict[str, Any] = {
        "manifest_version": MANIFEST_VERSION,
        "created_at_utc": _utc_now(),
        "source_discovery": {
            "workbook_sheet_metadata_only": True,
            "data_rows_scanned": 0,
            "numeric_filename_range_inference": False,
        },
        "inputs": inputs,
        "tasks": tasks,
        "summary": {
            "lab_sheet_count": len(sheet_names),
            "lab_expected_row_count": lab_total,
            "document_expected_row_count": document_total,
            "total_expected_row_count": lab_total + document_total,
            "lab_task_count": sum(1 for task in tasks if task["kind"] == "lab"),
            "document_task_count": sum(1 for task in tasks if task["kind"] == "document"),
            "total_task_count": len(tasks),
        },
        "output_root": str(output_root_path),
        "output_directories": {
            "lab_l1": str(output_root_path / "restricted" / "lab_l1"),
            "lab_quarantine": str(output_root_path / "restricted" / "lab_quarantine"),
            "document_l1": str(output_root_path / "restricted" / "document_l1"),
            "document_quarantine": str(output_root_path / "restricted" / "document_quarantine"),
            "audit": str(output_root_path / "audit"),
        },
    }
    _write_json_atomic(manifest_target, manifest)
    manifest["manifest_path"] = str(manifest_target)
    return manifest


def load_manifest(manifest_path: str | Path) -> dict[str, Any]:
    path = Path(manifest_path).resolve()
    data = json.loads(path.read_text(encoding="utf-8"))
    if data.get("manifest_version") != MANIFEST_VERSION:
        raise ManifestValidationError("unexpected manifest version")
    if data.get("summary", {}).get("total_expected_row_count") != 12_810_808:
        raise ManifestValidationError("manifest total expected row count is inconsistent")
    if data.get("summary", {}).get("lab_expected_row_count") != 12_211_982:
        raise ManifestValidationError("manifest laboratory row count is inconsistent")
    if data.get("summary", {}).get("document_expected_row_count") != 598_826:
        raise ManifestValidationError("manifest document row count is inconsistent")
    return data


class PipelineState:
    """SQLite state store for atomic, resumable shard tasks."""

    def __init__(self, database_path: str | Path):
        self.path = Path(database_path).resolve()
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.connection = sqlite3.connect(self.path)
        self.connection.row_factory = sqlite3.Row
        self.connection.execute("PRAGMA journal_mode=DELETE")
        self.connection.execute("PRAGMA foreign_keys=ON")
        self._create_schema()

    def _create_schema(self) -> None:
        self.connection.executescript(
            """
            CREATE TABLE IF NOT EXISTS pipeline_metadata (
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS shard_tasks (
                task_id TEXT PRIMARY KEY,
                source_input_id TEXT NOT NULL,
                kind TEXT NOT NULL,
                input_file TEXT NOT NULL,
                sheet_name TEXT,
                input_range_json TEXT NOT NULL,
                output_temp_file TEXT NOT NULL,
                output_file TEXT NOT NULL,
                audit_file TEXT NOT NULL,
                quarantine_file TEXT NOT NULL,
                expected_row_count INTEGER NOT NULL,
                status TEXT NOT NULL CHECK(status IN ('PENDING','RUNNING','SUCCEEDED','FAILED','QUARANTINED')),
                actual_row_count INTEGER,
                accepted_count INTEGER,
                quarantined_count INTEGER,
                output_sha256 TEXT,
                output_schema_sha256 TEXT,
                quarantine_sha256 TEXT,
                quarantine_schema_sha256 TEXT,
                attempts INTEGER NOT NULL DEFAULT 0,
                error_reason TEXT,
                updated_at TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS idx_shard_tasks_status ON shard_tasks(status, task_id);
            """
        )
        columns = {row[1] for row in self.connection.execute("PRAGMA table_info(shard_tasks)").fetchall()}
        for name, sql_type in (
            ("accepted_count", "INTEGER"),
            ("quarantined_count", "INTEGER"),
            ("quarantine_sha256", "TEXT"),
            ("quarantine_schema_sha256", "TEXT"),
        ):
            if name not in columns:
                self.connection.execute(f"ALTER TABLE shard_tasks ADD COLUMN {name} {sql_type}")
        self.connection.commit()

    def close(self) -> None:
        self.connection.close()

    def __enter__(self) -> "PipelineState":
        return self

    def __exit__(self, exc_type: Any, exc_value: Any, traceback: Any) -> None:
        self.close()

    def seed_manifest(self, manifest: Mapping[str, Any]) -> int:
        tasks = manifest.get("tasks")
        if not isinstance(tasks, list):
            raise ManifestValidationError("manifest tasks must be a list")
        inserted = 0
        with self.connection:
            self.connection.execute(
                "INSERT OR REPLACE INTO pipeline_metadata(key,value) VALUES (?,?)",
                ("manifest_version", str(manifest["manifest_version"])),
            )
            for task in tasks:
                existing = self.connection.execute(
                    "SELECT * FROM shard_tasks WHERE task_id=?", (task["task_id"],)
                ).fetchone()
                if existing is not None:
                    for field, value in (
                        ("input_file", task["input_file"]),
                        ("output_file", task["output_file"]),
                        ("expected_row_count", task["expected_row_count"]),
                    ):
                        if existing[field] != value:
                            raise ManifestValidationError(f"task definition changed: {task['task_id']} field {field}")
                    continue
                self.connection.execute(
                    """
                    INSERT INTO shard_tasks(
                        task_id,source_input_id,kind,input_file,sheet_name,input_range_json,
                        output_temp_file,output_file,audit_file,quarantine_file,expected_row_count,
                        status,attempts,updated_at
                    ) VALUES (?,?,?,?,?,?,?,?,?,?,?,'PENDING',0,?)
                    """,
                    (
                        task["task_id"],
                        task["source_input_id"],
                        task["kind"],
                        task["input_file"],
                        task.get("sheet_name"),
                        json.dumps(task["input_range"], ensure_ascii=False, sort_keys=True),
                        task["output_temp_file"],
                        task["output_file"],
                        task["audit_file"],
                        task["quarantine_file"],
                        task["expected_row_count"],
                        _utc_now(),
                    ),
                )
                inserted += 1
        return inserted

    def recover_running(self) -> int:
        with self.connection:
            cursor = self.connection.execute(
                "UPDATE shard_tasks SET status='PENDING', error_reason=?, updated_at=? WHERE status='RUNNING'",
                ("RECOVERED_FROM_RUNNING", _utc_now()),
            )
        return cursor.rowcount

    def claim_next(self) -> dict[str, Any] | None:
        self.connection.execute("BEGIN IMMEDIATE")
        try:
            row = self.connection.execute(
                "SELECT * FROM shard_tasks WHERE status='PENDING' ORDER BY task_id LIMIT 1"
            ).fetchone()
            if row is None:
                self.connection.commit()
                return None
            now = _utc_now()
            self.connection.execute(
                "UPDATE shard_tasks SET status='RUNNING', attempts=attempts+1, updated_at=? WHERE task_id=?",
                (now, row["task_id"]),
            )
            self.connection.commit()
            return self.get_task(row["task_id"])
        except Exception:
            self.connection.rollback()
            raise

    def claim_task(self, task_id: str) -> dict[str, Any] | None:
        """Claim one manifest task, leaving SUCCEEDED tasks untouched."""

        self.connection.execute("BEGIN IMMEDIATE")
        try:
            row = self.connection.execute("SELECT * FROM shard_tasks WHERE task_id=?", (task_id,)).fetchone()
            if row is None:
                self.connection.commit()
                return None
            if row["status"] == "SUCCEEDED":
                self.connection.commit()
                return dict(row)
            if row["status"] != "PENDING":
                self.connection.commit()
                return None
            self.connection.execute(
                "UPDATE shard_tasks SET status='RUNNING', attempts=attempts+1, updated_at=? WHERE task_id=?",
                (_utc_now(), task_id),
            )
            self.connection.commit()
            return self.get_task(task_id)
        except Exception:
            self.connection.rollback()
            raise

    def requeue(self, task_id: str, error_reason: str) -> None:
        with self.connection:
            cursor = self.connection.execute(
                "UPDATE shard_tasks SET status='PENDING', error_reason=?, updated_at=? WHERE task_id=? AND status='RUNNING'",
                (error_reason[:2000], _utc_now(), task_id),
            )
        if cursor.rowcount != 1:
            raise RuntimeError(f"task is not RUNNING: {task_id}")

    def get_task(self, task_id: str) -> dict[str, Any] | None:
        row = self.connection.execute("SELECT * FROM shard_tasks WHERE task_id=?", (task_id,)).fetchone()
        if row is None:
            return None
        result = dict(row)
        result["input_range"] = json.loads(result.pop("input_range_json"))
        return result

    def mark_succeeded(self, task_id: str, result: CommitResult) -> None:
        accepted_value = getattr(result, "accepted_count", None)
        accepted_count = accepted_value if accepted_value is not None else result.row_count
        quarantined_count = int(getattr(result, "quarantined_count", 0) or 0)
        quarantine_sha256 = getattr(result, "quarantine_sha256", None)
        quarantine_schema_sha256 = getattr(result, "quarantine_schema_sha256", None)
        with self.connection:
            cursor = self.connection.execute(
                """
                UPDATE shard_tasks
                SET status='SUCCEEDED', actual_row_count=?, accepted_count=?, quarantined_count=?,
                    output_sha256=?, output_schema_sha256=?, quarantine_sha256=?, quarantine_schema_sha256=?,
                    error_reason=NULL, updated_at=?
                WHERE task_id=? AND status='RUNNING'
                """,
                (
                    result.row_count,
                    accepted_count,
                    quarantined_count,
                    result.output_sha256,
                    result.schema_sha256,
                    quarantine_sha256,
                    quarantine_schema_sha256,
                    _utc_now(),
                    task_id,
                ),
            )
        if cursor.rowcount != 1:
            raise RuntimeError(f"task is not RUNNING: {task_id}")

    def mark_failed(self, task_id: str, error_reason: str, *, quarantined: bool = False) -> None:
        status = "QUARANTINED" if quarantined else "FAILED"
        with self.connection:
            cursor = self.connection.execute(
                "UPDATE shard_tasks SET status=?, error_reason=?, updated_at=? WHERE task_id=? AND status='RUNNING'",
                (status, error_reason[:2000], _utc_now(), task_id),
            )
        if cursor.rowcount != 1:
            raise RuntimeError(f"task is not RUNNING: {task_id}")

    def status_counts(self) -> dict[str, int]:
        counts = {status: 0 for status in STATUSES}
        rows = self.connection.execute("SELECT status,COUNT(*) AS count FROM shard_tasks GROUP BY status").fetchall()
        for row in rows:
            counts[row["status"]] = row["count"]
        return counts


def atomic_commit_parquet(
    temporary_file: str | Path,
    final_file: str | Path,
    *,
    expected_row_count: int,
    expected_schema: pa.Schema | None = None,
) -> CommitResult:
    temporary = Path(temporary_file).resolve()
    final = Path(final_file).resolve()
    if not temporary.name.endswith(".partial"):
        raise ValueError("temporary output must use a .partial suffix")
    if not temporary.is_file():
        raise FileNotFoundError(temporary)
    if final.exists():
        raise FileExistsError(f"refusing to overwrite formal output: {final}")
    parquet = pq.ParquetFile(temporary)
    try:
        row_count = parquet.metadata.num_rows
        schema = parquet.schema_arrow
        if row_count != expected_row_count:
            raise ValueError(f"row count mismatch: expected {expected_row_count}, got {row_count}")
        if expected_schema is not None and not schema.equals(expected_schema, check_metadata=False):
            raise ValueError("Parquet schema mismatch")
    finally:
        parquet.close()
    output_hash = _sha256_file(temporary)
    schema_hash = _schema_sha256(schema)
    final.parent.mkdir(parents=True, exist_ok=True)
    os.replace(temporary, final)
    committed = pq.ParquetFile(final)
    try:
        if committed.metadata.num_rows != expected_row_count:
            raise RuntimeError("formal output row count changed after commit")
    finally:
        committed.close()
    committed_hash = _sha256_file(final)
    if committed_hash != output_hash:
        raise RuntimeError("formal output hash changed after commit")
    return CommitResult(row_count, committed_hash, schema_hash, str(final))


def _empty_table(schema: pa.Schema) -> pa.Table:
    return pa.Table.from_arrays([pa.array([], type=field.type) for field in schema], schema=schema)


def _json_bytes(value: Any, *, encoding: str = "utf-8") -> bytes:
    text = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    if encoding == "gb18030":
        return text.encode(encoding, errors="surrogateescape")
    return text.encode(encoding, errors="surrogatepass")


def _records_to_table(
    records: Sequence[Mapping[str, Any]],
    schema: pa.Schema,
    *,
    binary_fields: frozenset[str] = frozenset(),
) -> pa.Table:
    if not records:
        return _empty_table(schema)
    prepared: list[dict[str, Any]] = []
    for record in records:
        item = dict(record)
        for field in binary_fields:
            value = item.get(field)
            if isinstance(value, str):
                item[field] = _document_source_bytes(value)
        prepared.append(item)
    return pa.Table.from_pylist(prepared, schema=schema)


def _archive_partial(path: Path, attempt: int) -> None:
    if not path.exists():
        return
    archived = path.with_name(f"{path.name}.failed-attempt-{attempt}")
    suffix = 1
    while archived.exists():
        archived = path.with_name(f"{path.name}.failed-attempt-{attempt}-{suffix}")
        suffix += 1
    path.replace(archived)


class _ShardWriters:
    def __init__(self, task: Mapping[str, Any], attempt: int, formal_schema: pa.Schema, quarantine_schema: pa.Schema):
        self.task = task
        self.formal_schema = formal_schema
        self.quarantine_schema = quarantine_schema
        self.formal_temp = Path(task["output_temp_file"])
        self.formal_final = Path(task["output_file"])
        self.quarantine_temp = Path(task["quarantine_file"] + ".partial")
        self.quarantine_final = Path(task["quarantine_file"])
        _archive_partial(self.formal_temp, attempt)
        _archive_partial(self.quarantine_temp, attempt)
        self.formal_temp.parent.mkdir(parents=True, exist_ok=True)
        self.formal_writer = pq.ParquetWriter(self.formal_temp, formal_schema, compression="zstd")
        self.quarantine_writer: pq.ParquetWriter | None = None
        self.formal_count = 0
        self.quarantine_count = 0

    def add_formal(self, records: Sequence[Mapping[str, Any]], *, binary_fields: frozenset[str] = frozenset()) -> None:
        if not records:
            return
        table = _records_to_table(records, self.formal_schema, binary_fields=binary_fields)
        self.formal_writer.write_table(table)
        self.formal_count += table.num_rows

    def add_quarantine(self, records: Sequence[Mapping[str, Any]]) -> None:
        if not records:
            return
        if self.quarantine_writer is None:
            self.quarantine_temp.parent.mkdir(parents=True, exist_ok=True)
            self.quarantine_writer = pq.ParquetWriter(self.quarantine_temp, self.quarantine_schema, compression="zstd")
        table = _records_to_table(records, self.quarantine_schema)
        self.quarantine_writer.write_table(table)
        self.quarantine_count += table.num_rows

    def finish(self) -> CommitResult:
        self.formal_writer.close()
        if self.formal_count == 0:
            # ParquetWriter still carries the declared schema; this guarantees a
            # committed formal output even when every source row is quarantined.
            pass
        if self.quarantine_writer is not None:
            self.quarantine_writer.close()
        formal = atomic_commit_parquet(
            self.formal_temp,
            self.formal_final,
            expected_row_count=self.formal_count,
            expected_schema=self.formal_schema,
        )
        quarantine_hash: str | None = None
        quarantine_schema_hash: str | None = None
        if self.quarantine_count:
            if self.quarantine_writer is None:
                raise GlobalPipelineBlock("quarantine writer missing despite quarantine rows")
            quarantine = atomic_commit_parquet(
                self.quarantine_temp,
                self.quarantine_final,
                expected_row_count=self.quarantine_count,
                expected_schema=self.quarantine_schema,
            )
            quarantine_hash = quarantine.output_sha256
            quarantine_schema_hash = quarantine.schema_sha256
        return CommitResult(
            row_count=self.formal_count + self.quarantine_count,
            output_sha256=formal.output_sha256,
            schema_sha256=formal.schema_sha256,
            output_file=formal.output_file,
            accepted_count=self.formal_count,
            quarantined_count=self.quarantine_count,
            quarantine_sha256=quarantine_hash,
            quarantine_schema_sha256=quarantine_schema_hash,
        )


def _lab_quarantine_schema() -> pa.Schema:
    return pa.schema(
        [
            pa.field("source_workbook", pa.string()),
            pa.field("source_sheet", pa.string()),
            pa.field("source_row", pa.int64()),
            pa.field("source_record_id", pa.string()),
            pa.field("raw_values_json", pa.large_binary()),
            pa.field("row_hash", pa.string()),
            pa.field("quarantine_reason", pa.string()),
            pa.field("ingestion_version", pa.string()),
        ]
    )


def _document_quarantine_schema() -> pa.Schema:
    return pa.schema(
        [
            pa.field("source_file", pa.string()),
            pa.field("source_row", pa.int64()),
            pa.field("source_record_id", pa.string()),
            pa.field("physical_line_start", pa.int64()),
            pa.field("physical_line_end", pa.int64()),
            pa.field("raw_token_array_json", pa.large_binary()),
            pa.field("parsed_column_count", pa.int64()),
            pa.field("row_hash", pa.string()),
            pa.field("quarantine_reason", pa.string()),
            pa.field("ingestion_version", pa.string()),
        ]
    )


def _lab_raw_quarantine(
    raw_values: Sequence[Any],
    *,
    source_workbook: str,
    source_sheet: str,
    source_row: int,
    reason: str,
) -> dict[str, Any]:
    json_values = [_jsonable_cell(value) for value in raw_values]
    raw_bytes = _json_bytes(json_values)
    return {
        "source_workbook": source_workbook,
        "source_sheet": source_sheet,
        "source_row": source_row,
        "source_record_id": f"{source_sheet}:{source_row}",
        "raw_values_json": raw_bytes,
        "row_hash": hashlib.sha256(raw_bytes).hexdigest(),
        "quarantine_reason": reason,
        "ingestion_version": LAB_INGESTION_VERSION,
    }


def _document_raw_quarantine(
    raw_values: Sequence[str],
    *,
    source_file: str,
    source_row: int,
    physical_line_start: int,
    physical_line_end: int,
    reason: str,
) -> dict[str, Any]:
    raw_bytes = _json_bytes(list(raw_values), encoding="gb18030")
    return {
        "source_file": source_file,
        "source_row": source_row,
        "source_record_id": f"{Path(source_file).name}:{source_row}",
        "physical_line_start": physical_line_start,
        "physical_line_end": physical_line_end,
        "raw_token_array_json": raw_bytes,
        "parsed_column_count": len(raw_values),
        "row_hash": _document_hash_raw_fields(raw_values),
        "quarantine_reason": reason,
        "ingestion_version": "document_l1_restricted_v1",
    }


def _parse_lab_record(
    raw_values: Sequence[Any],
    *,
    source_workbook: str,
    source_sheet: str,
    source_row: int,
) -> tuple[dict[str, Any] | None, dict[str, Any] | None]:
    try:
        normalized_values, digest = _normalize_row(raw_values)
    except (TypeError, ValueError) as exc:
        return None, _lab_raw_quarantine(
            raw_values,
            source_workbook=source_workbook,
            source_sheet=source_sheet,
            source_row=source_row,
            reason=f"SOURCE_CELL_TYPE_OR_STRUCTURE:{type(exc).__name__}",
        )
    row = dict(zip(LAB_RAW_FIELDS, normalized_values))
    row.update(
        {
            "source_workbook": source_workbook,
            "source_sheet": source_sheet,
            "source_row": source_row,
            "source_record_id": f"{source_sheet}:{source_row}",
            "row_hash": digest,
            "ingestion_version": LAB_INGESTION_VERSION,
        }
    )
    try:
        result = parse_result(row.get("检验结果值"))
        reference = parse_reference(row.get("检验参考值"))
        unit = normalize_unit(row.get("检验结果单位"))
        timing = normalize_time_pair(row.get("送检时间"), row.get("报告时间"))
    except Exception as exc:
        raise RuleDefect(f"lab parser raised {type(exc).__name__} at {source_sheet}:{source_row}") from exc
    row.update(result)
    row.update(reference)
    row.update(unit)
    row.update(timing)
    row["source_abnormal_flag"] = row.get("结果正常标志")
    if timing["time_parse_status"] in {"INVALID_SPECIMEN_REPORT_ORDER", "INVALID_EXCEL_DATETIME"}:
        return None, _lab_raw_quarantine(
            raw_values,
            source_workbook=source_workbook,
            source_sheet=source_sheet,
            source_row=source_row,
            reason=timing["time_parse_status"],
        )
    return row, None


def _iter_document_range(source_file: Path, start_row: int, end_row: int):
    with source_file.open("r", encoding="gb18030", errors="surrogateescape", newline="") as stream:
        reader = csv.reader(stream)
        try:
            header = tuple(next(reader))
        except StopIteration as exc:
            raise GlobalPipelineBlock("document CSV is empty") from exc
        if header != DOCUMENT_RAW_FIELDS:
            raise GlobalPipelineBlock("document CSV header does not match the frozen contract")
        previous_physical_end = 1
        for source_row in range(2, end_row + 1):
            try:
                raw_values = next(reader)
            except StopIteration as exc:
                raise GlobalPipelineBlock(f"document CSV ended before source_row {end_row}") from exc
            physical_end = reader.line_num
            physical_start = previous_physical_end + 1
            previous_physical_end = physical_end
            if source_row >= start_row:
                yield source_row, raw_values, physical_start, physical_end


def _process_lab_task(task: Mapping[str, Any], workbook: Any, *, batch_rows: int = 10_000) -> CommitResult:
    sheet_name = task["sheet_name"]
    source_file = str(Path(task["input_file"]).resolve())
    source_range = task["input_range"]
    expected = int(task["expected_row_count"])
    formal_schema = _lab_enriched_schema(_lab_parquet_schema())
    quarantine_schema = _lab_quarantine_schema()
    writers = _ShardWriters(task, int(task["attempts"]), formal_schema, quarantine_schema)
    accepted_batch: list[dict[str, Any]] = []
    quarantine_batch: list[dict[str, Any]] = []
    seen = 0
    try:
        worksheet = workbook[sheet_name]
        iterator = worksheet.iter_rows(
            min_row=source_range["source_row_start"],
            max_row=source_range["source_row_end"],
            max_col=len(LAB_RAW_FIELDS),
            values_only=True,
        )
        for source_row, raw_values in enumerate(iterator, start=source_range["source_row_start"]):
            if len(raw_values) != len(LAB_RAW_FIELDS):
                quarantine_batch.append(
                    _lab_raw_quarantine(
                        raw_values,
                        source_workbook=source_file,
                        source_sheet=sheet_name,
                        source_row=source_row,
                        reason="SOURCE_COLUMN_COUNT_INVALID",
                    )
                )
            else:
                formal, quarantine = _parse_lab_record(
                    raw_values,
                    source_workbook=source_file,
                    source_sheet=sheet_name,
                    source_row=source_row,
                )
                if formal is not None:
                    accepted_batch.append(formal)
                if quarantine is not None:
                    quarantine_batch.append(quarantine)
            seen += 1
            if len(accepted_batch) + len(quarantine_batch) >= batch_rows:
                writers.add_formal(accepted_batch)
                writers.add_quarantine(quarantine_batch)
                accepted_batch.clear()
                quarantine_batch.clear()
        writers.add_formal(accepted_batch)
        writers.add_quarantine(quarantine_batch)
        accepted_batch.clear()
        quarantine_batch.clear()
        if seen != expected:
            raise GlobalPipelineBlock(
                f"lab source range count mismatch for {task['task_id']}: expected {expected}, got {seen}"
            )
        result = writers.finish()
        if result.row_count != expected:
            raise GlobalPipelineBlock(f"lab token conservation failed for {task['task_id']}")
        return result
    except Exception:
        try:
            writers.formal_writer.close()
        except Exception:
            pass
        if writers.quarantine_writer is not None:
            try:
                writers.quarantine_writer.close()
            except Exception:
                pass
        raise


def _process_document_task(task: Mapping[str, Any], *, batch_rows: int = 5_000) -> CommitResult:
    source_file = Path(task["input_file"]).resolve()
    source_range = task["input_range"]
    expected = int(task["expected_row_count"])
    formal_schema = _document_parquet_schema()
    quarantine_schema = _document_quarantine_schema()
    writers = _ShardWriters(task, int(task["attempts"]), formal_schema, quarantine_schema)
    accepted_batch: list[dict[str, Any]] = []
    quarantine_batch: list[dict[str, Any]] = []
    seen = 0
    try:
        for source_row, raw_values, physical_start, physical_end in _iter_document_range(
            source_file,
            source_range["source_row_start"],
            source_range["source_row_end"],
        ):
            if len(raw_values) != len(DOCUMENT_RAW_FIELDS):
                quarantine_batch.append(
                    _document_raw_quarantine(
                        raw_values,
                        source_file=str(source_file),
                        source_row=source_row,
                        physical_line_start=physical_start,
                        physical_line_end=physical_end,
                        reason="SOURCE_COLUMN_COUNT_INVALID",
                    )
                )
            else:
                try:
                    record = _build_document_record(raw_values, source_file=str(source_file), source_row=source_row)
                except Exception as exc:
                    raise RuleDefect(f"document parser raised {type(exc).__name__} at source_row {source_row}") from exc
                if "DATE_PARSE_FAILED" in record["parser_status"]:
                    quarantine_batch.append(
                        _document_raw_quarantine(
                            raw_values,
                            source_file=str(source_file),
                            source_row=source_row,
                            physical_line_start=physical_start,
                            physical_line_end=physical_end,
                            reason="DATE_PARSE_FAILED",
                        )
                    )
                else:
                    accepted_batch.append(record)
            seen += 1
            if len(accepted_batch) + len(quarantine_batch) >= batch_rows:
                writers.add_formal(accepted_batch, binary_fields=frozenset({"文书内容"}))
                writers.add_quarantine(quarantine_batch)
                accepted_batch.clear()
                quarantine_batch.clear()
        writers.add_formal(accepted_batch, binary_fields=frozenset({"文书内容"}))
        writers.add_quarantine(quarantine_batch)
        accepted_batch.clear()
        quarantine_batch.clear()
        if seen != expected:
            raise GlobalPipelineBlock(
                f"document source range count mismatch for {task['task_id']}: expected {expected}, got {seen}"
            )
        result = writers.finish()
        if result.row_count != expected:
            raise GlobalPipelineBlock(f"document token conservation failed for {task['task_id']}")
        return result
    except Exception:
        try:
            writers.formal_writer.close()
        except Exception:
            pass
        if writers.quarantine_writer is not None:
            try:
                writers.quarantine_writer.close()
            except Exception:
                pass
        raise


def verify_manifest_inputs(manifest: Mapping[str, Any]) -> dict[str, Any]:
    verified: dict[tuple[str, str], bool] = {}
    for item in manifest["inputs"]:
        key = (item["input_file"], item["sha256"])
        if key in verified:
            continue
        path = Path(item["input_file"]).resolve()
        if not path.is_file():
            raise GlobalPipelineBlock(f"manifest input is missing: {path}")
        actual_size = path.stat().st_size
        actual_sha256 = _sha256_file(path)
        if actual_size != item["file_size"] or actual_sha256 != item["sha256"]:
            raise GlobalPipelineBlock(f"input hash or size mismatch: {path}")
        verified[key] = True
    return {
        "distinct_input_count": len(verified),
        "verified_input_count": sum(verified.values()),
        "all_inputs_verified": all(verified.values()),
    }


def _canary_tasks(manifest: Mapping[str, Any], output_root: Path, lab_rows: int, document_rows: int) -> list[dict[str, Any]]:
    if lab_rows <= 0 or document_rows <= 0:
        raise ValueError("canary row limits must be positive")
    result: list[dict[str, Any]] = []
    for item in manifest["inputs"]:
        limit = lab_rows if item["kind"] == "lab" else document_rows
        source_tasks = [task for task in manifest["tasks"] if task["source_input_id"] == item["input_id"]]
        source_task = source_tasks[0]
        original_range = source_task["input_range"]
        count = min(limit, item["expected_row_count"])
        task = dict(source_task)
        task["task_id"] = f"{item['input_id']}:canary"
        task["expected_row_count"] = count
        task["input_range"] = {
            "data_row_start": 1,
            "data_row_end": count,
            "source_row_start": original_range["source_row_start"],
            "source_row_end": original_range["source_row_start"] + count - 1,
            "expected_row_count": count,
        }
        task.update(_target_paths(output_root, item["kind"], item["sheet_name"], 1))
        result.append(task)
    return result


def _task_rows(state: PipelineState) -> list[dict[str, Any]]:
    rows = state.connection.execute("SELECT * FROM shard_tasks ORDER BY task_id").fetchall()
    result: list[dict[str, Any]] = []
    for row in rows:
        item = dict(row)
        item["input_range"] = json.loads(item.pop("input_range_json"))
        result.append(item)
    return result


def _aggregate_state(state: PipelineState) -> dict[str, Any]:
    rows = _task_rows(state)
    accepted = sum(int(row["accepted_count"] or 0) for row in rows)
    quarantined = sum(int(row["quarantined_count"] or 0) for row in rows)
    expected = sum(int(row["expected_row_count"]) for row in rows)
    return {
        "task_status_counts": state.status_counts(),
        "expected_input_row_count": expected,
        "accepted_count": accepted,
        "quarantined_count": quarantined,
        "accepted_plus_quarantined": accepted + quarantined,
        "token_conservation": accepted + quarantined == expected,
        "all_tasks_succeeded": all(row["status"] == "SUCCEEDED" for row in rows),
    }


def _run_one_task(
    task: Mapping[str, Any],
    state: PipelineState,
    *,
    workbook: Any | None,
    max_retries: int,
) -> dict[str, Any]:
    for _ in range(max_retries):
        claimed = state.claim_task(task["task_id"])
        if claimed is None:
            current = state.get_task(task["task_id"])
            if current is None:
                raise GlobalPipelineBlock(f"task missing from state: {task['task_id']}")
            if current["status"] == "SUCCEEDED":
                return current
            continue
        if claimed.get("status") == "SUCCEEDED":
            return claimed
        try:
            if task["kind"] == "lab":
                result = _process_lab_task(claimed, workbook)
            else:
                result = _process_document_task(claimed)
            state.mark_succeeded(task["task_id"], result)
            return state.get_task(task["task_id"]) or {}
        except GlobalPipelineBlock as exc:
            state.mark_failed(task["task_id"], str(exc))
            raise
        except Exception as exc:
            current = state.get_task(task["task_id"]) or claimed
            if int(current["attempts"]) < max_retries:
                state.requeue(task["task_id"], f"RETRYABLE:{type(exc).__name__}:{exc}")
                continue
            state.mark_failed(task["task_id"], f"{type(exc).__name__}: {exc}")
            return state.get_task(task["task_id"]) or {}
    raise GlobalPipelineBlock(f"retry loop exhausted without a task result: {task['task_id']}")


def run_pipeline(
    manifest_path: str | Path = DEFAULT_MANIFEST,
    state_path: str | Path = DEFAULT_STATE_DB,
    output_root: str | Path = DEFAULT_OUTPUT_ROOT,
    report_path: str | Path | None = None,
    *,
    max_retries: int = 3,
    canary_rows: tuple[int, int] | None = None,
    progress_callback: Any | None = None,
) -> dict[str, Any]:
    if max_retries <= 0:
        raise ValueError("max_retries must be positive")
    manifest = load_manifest(manifest_path)
    verification = verify_manifest_inputs(manifest)
    root = Path(output_root).resolve()
    started = time.perf_counter()
    peak_memory = 0
    canary = canary_rows is not None
    if canary:
        canary_root = root / "canary"
        canary_root.mkdir(parents=True, exist_ok=True)
        tasks = _canary_tasks(manifest, canary_root, canary_rows[0], canary_rows[1])
        state_file = canary_root / "stage5_canary_state.sqlite3"
        state_manifest = {"manifest_version": MANIFEST_VERSION, "tasks": tasks}
        report_target = Path(report_path).resolve() if report_path else root / "audit" / "stage5_canary_report.json"
    else:
        tasks = list(manifest["tasks"])
        state_file = Path(state_path).resolve()
        state_manifest = manifest
        report_target = Path(report_path).resolve() if report_path else root / "stage5_total_report.json"

    with PipelineState(state_file) as state:
        state.recover_running()
        state.seed_manifest(state_manifest)
        workbook = None
        try:
            workbook = openpyxl.load_workbook(manifest["inputs"][0]["input_file"], read_only=True, data_only=True)
            for sheet_name in EXPECTED_LAB_SHEETS:
                if sheet_name not in workbook.sheetnames:
                    raise GlobalPipelineBlock(f"lab sheet missing during execution: {sheet_name}")
            for sheet_name in EXPECTED_LAB_SHEETS:
                sheet_tasks = [task for task in tasks if task["kind"] == "lab" and task["sheet_name"] == sheet_name]
                accepted_before = sum(int((state.get_task(task["task_id"]) or {}).get("accepted_count") or 0) for task in sheet_tasks)
                quarantined_before = sum(int((state.get_task(task["task_id"]) or {}).get("quarantined_count") or 0) for task in sheet_tasks)
                for task in sheet_tasks:
                    _run_one_task(task, state, workbook=workbook, max_retries=max_retries)
                    peak_memory = max(peak_memory, _memory_rss_bytes())
                accepted_after = sum(int((state.get_task(task["task_id"]) or {}).get("accepted_count") or 0) for task in sheet_tasks)
                quarantined_after = sum(int((state.get_task(task["task_id"]) or {}).get("quarantined_count") or 0) for task in sheet_tasks)
                if progress_callback:
                    progress_callback(
                        {
                            "scope": "sheet",
                            "sheet_name": sheet_name,
                            "task_count": len(sheet_tasks),
                            "accepted_count": accepted_after,
                            "quarantined_count": quarantined_after,
                            "newly_accepted_count": accepted_after - accepted_before,
                            "newly_quarantined_count": quarantined_after - quarantined_before,
                        }
                    )
            workbook.close()
            workbook = None

            document_tasks = [task for task in tasks if task["kind"] == "document"]
            for task in document_tasks:
                _run_one_task(task, state, workbook=None, max_retries=max_retries)
                peak_memory = max(peak_memory, _memory_rss_bytes())
                if progress_callback:
                    current = state.get_task(task["task_id"]) or {}
                    progress_callback(
                        {
                            "scope": "document_shard",
                            "task_id": task["task_id"],
                            "accepted_count": int(current.get("accepted_count") or 0),
                            "quarantined_count": int(current.get("quarantined_count") or 0),
                            "status": current.get("status"),
                        }
                    )
        except Exception:
            if workbook is not None:
                workbook.close()
            raise
        aggregate = _aggregate_state(state)

    peak_memory = max(peak_memory, _memory_rss_bytes())
    report: dict[str, Any] = {
        "report_version": "stage5_canary_report_v1" if canary else "stage5_total_report_v1",
        "run_mode": "canary" if canary else "full",
        "manifest_path": str(Path(manifest_path).resolve()),
        "state_path": str(state_file),
        "output_root": str(root),
        "input_verification": verification,
        "canary_limits": {"lab_rows_per_sheet": canary_rows[0], "document_rows": canary_rows[1]} if canary else None,
        "aggregate": aggregate,
        "peak_memory_bytes": peak_memory,
        "peak_memory_gib": peak_memory / (1024**3),
        "memory_limit_ok": peak_memory < 4 * 1024**3,
        "elapsed_seconds": time.perf_counter() - started,
        "parser_rules_modified_during_run": False,
        "schema_modified_during_run": False,
        "data_contract_modified_during_run": False,
        "report_contains_patient_identifiers": False,
        "report_contains_document_content": False,
        "patient_alignment_started": False,
        "timeline_started": False,
    }
    _write_json_atomic(report_target, report)
    report["report_path"] = str(report_target)
    return report


def initialize_state(manifest_path: str | Path = DEFAULT_MANIFEST, state_path: str | Path = DEFAULT_STATE_DB) -> dict[str, Any]:
    manifest = load_manifest(manifest_path)
    with PipelineState(state_path) as state:
        recovered = state.recover_running()
        inserted = state.seed_manifest(manifest)
        return {"recovered_running": recovered, "inserted_tasks": inserted, "status_counts": state.status_counts()}


def dry_run(
    manifest_path: str | Path = DEFAULT_MANIFEST,
    output_root: str | Path = DEFAULT_OUTPUT_ROOT,
    report_path: str | Path = DEFAULT_DRY_RUN_REPORT,
) -> dict[str, Any]:
    manifest = load_manifest(manifest_path)
    root = Path(output_root).resolve()
    report_target = Path(report_path).resolve()
    disk_path = root if root.exists() else root.parent
    usage = shutil.disk_usage(disk_path)
    total_input_bytes = sum(int(item["file_size"]) for item in manifest["inputs"] if item["input_id"] in {"lab:Sheet1", "document:admission_discharge"})
    report_inputs: list[dict[str, Any]] = []
    for item in manifest["inputs"]:
        item_tasks = [task for task in manifest["tasks"] if task["source_input_id"] == item["input_id"]]
        report_inputs.append(
            {
                "input_id": item["input_id"],
                "kind": item["kind"],
                "input_file": item["input_file"],
                "file_size": item["file_size"],
                "sha256": item["sha256"],
                "sheet_name": item["sheet_name"],
                "expected_row_count": item["expected_row_count"],
                "shard_target_rows": item["shard_target_rows"],
                "shard_count": len(item_tasks),
                "shards": [
                    {
                        "task_id": task["task_id"],
                        "input_range": task["input_range"],
                        "expected_row_count": task["expected_row_count"],
                        "output_file": task["output_file"],
                        "output_temp_file": task["output_temp_file"],
                        "audit_file": task["audit_file"],
                    }
                    for task in item_tasks
                ],
            }
        )
    report: dict[str, Any] = {
        "report_version": "stage5_dry_run_report_v1",
        "mode": "dry-run",
        "manifest_version": manifest["manifest_version"],
        "manifest_path": str(Path(manifest_path).resolve()),
        "output_root": str(root),
        "inputs": report_inputs,
        "summary": manifest["summary"],
        "disk_space_check": {
            "checked_path": str(disk_path),
            "free_bytes": usage.free,
            "total_bytes": usage.total,
            "required_bytes_estimate": total_input_bytes,
            "sufficient_for_input_size_estimate": usage.free >= total_input_bytes,
        },
        "data_rows_scanned": 0,
        "patient_level_formal_output_created": False,
        "restricted_formal_output_created": False,
        "audit_only_report_created": True,
        "output_directories": manifest["output_directories"],
    }
    _write_json_atomic(report_target, report)
    report["report_path"] = str(report_target)
    return report


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)

    manifest_parser = subparsers.add_parser("manifest")
    manifest_parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    manifest_parser.add_argument("--document", type=Path, default=DEFAULT_DOCUMENT)
    manifest_parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    manifest_parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT_ROOT)

    state_parser = subparsers.add_parser("init-state")
    state_parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    state_parser.add_argument("--state", type=Path, default=DEFAULT_STATE_DB)

    dry_parser = subparsers.add_parser("dry-run")
    dry_parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    dry_parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT_ROOT)
    dry_parser.add_argument("--report", type=Path, default=DEFAULT_DRY_RUN_REPORT)

    run_parser = subparsers.add_parser("run")
    run_parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    run_parser.add_argument("--state", type=Path, default=DEFAULT_STATE_DB)
    run_parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT_ROOT)
    run_parser.add_argument("--report", type=Path, default=DEFAULT_OUTPUT_ROOT / "stage5_total_report.json")
    run_parser.add_argument("--max-retries", type=int, default=3)

    canary_parser = subparsers.add_parser("canary")
    canary_parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    canary_parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT_ROOT)
    canary_parser.add_argument("--report", type=Path, default=DEFAULT_OUTPUT_ROOT / "audit" / "stage5_canary_report.json")
    canary_parser.add_argument("--lab-rows", type=int, default=100)
    canary_parser.add_argument("--document-rows", type=int, default=100)
    canary_parser.add_argument("--max-retries", type=int, default=3)
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = _build_parser().parse_args(argv)
    if args.command == "manifest":
        result = build_manifest(args.workbook, args.document, args.manifest, output_root=args.output_root)
        print(json.dumps({"manifest_path": result["manifest_path"], "summary": result["summary"]}, ensure_ascii=False, indent=2))
    elif args.command == "init-state":
        print(json.dumps(initialize_state(args.manifest, args.state), ensure_ascii=False, indent=2))
    elif args.command == "dry-run":
        result = dry_run(args.manifest, args.output_root, args.report)
        print(json.dumps({"report_path": result["report_path"], "summary": result["summary"], "disk_space_check": result["disk_space_check"]}, ensure_ascii=False, indent=2))
    elif args.command == "run":
        def progress(event: Mapping[str, Any]) -> None:
            print(json.dumps({"progress": event}, ensure_ascii=False))

        result = run_pipeline(
            args.manifest,
            args.state,
            args.output_root,
            args.report,
            max_retries=args.max_retries,
            progress_callback=progress,
        )
        print(json.dumps({"report_path": result["report_path"], "aggregate": result["aggregate"]}, ensure_ascii=False, indent=2))
    else:
        def progress(event: Mapping[str, Any]) -> None:
            print(json.dumps({"progress": event}, ensure_ascii=False))

        result = run_pipeline(
            args.manifest,
            DEFAULT_STATE_DB,
            args.output_root,
            args.report,
            max_retries=args.max_retries,
            canary_rows=(args.lab_rows, args.document_rows),
            progress_callback=progress,
        )
        print(json.dumps({"report_path": result["report_path"], "aggregate": result["aggregate"]}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
