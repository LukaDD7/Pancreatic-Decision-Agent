"""Stage 6 pathology increment layer.

This module is deliberately independent from ``stage6_alignment_v2.py``.
The formal Stage 6 state and outputs are opened read-only; pathology output
is written to an isolated increment directory and state database.
"""

from __future__ import annotations

import argparse
from contextlib import closing
import hashlib
import json
import math
import os
import re
import sqlite3
import tempfile
import uuid
from collections import OrderedDict
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook

from data_pipeline.paths import data_root


DATA_ROOT = data_root()
PATHOLOGY_ROOT = DATA_ROOT / "病理报告"
STAGE6_ROOT = DATA_ROOT / "pipeline_outputs_stage6_v2"
STAGE6_RESTRICTED = STAGE6_ROOT / "restricted"
STAGE6_FULL_ROOT = STAGE6_RESTRICTED / "full"
STAGE6_STATE = STAGE6_RESTRICTED / "state" / "stage6_full_state_v2.sqlite3"
INCREMENT_ROOT = STAGE6_RESTRICTED / "increment" / "pathology_v1"
INCREMENT_FULL_ROOT = INCREMENT_ROOT / "full"
INCREMENT_CANARY_ROOT = INCREMENT_ROOT / "canary_v2"
INCREMENT_STATE = STAGE6_RESTRICTED / "state" / "stage6_pathology_increment_v1.sqlite3"
MANIFEST_PATH = INCREMENT_ROOT / "pathology_manifest_v1.json"
PREFLIGHT_REPORT = INCREMENT_ROOT / "pathology_preflight_report_v1.json"
CANARY_REPORT = INCREMENT_ROOT / "pathology_canary_report_v2.json"
FULL_REPORT = INCREMENT_ROOT / "pathology_increment_report_v1.json"
VALIDATION_REPORT = INCREMENT_ROOT / "pathology_validation_report_v1.json"

RULE_VERSION = "stage6_pathology_increment_rules_v1"
INGESTION_VERSION = "stage6_pathology_increment_v1"
SEED = "stage6-pathology-canary-v1"
PATIENT_NAMESPACE = uuid.UUID("6c5b2c26-80b6-5b13-9aa3-7c9b3b6b6c1d")
PATHOLOGY_NAMESPACE = uuid.UUID("f3c8ad36-a926-5b5e-93e3-0f10f5bc1d1a")
ID_CARD_RE = re.compile(r"^\d{17}[0-9Xx]$")
MISSING = {"", "nan", "none", "null", "nat", "<na>"}

DISEASE_INPUTS = (
    ("PDAC", "01_PDAC"),
    ("IPMN", "02_IPMN"),
    ("PanNET/PNET", "03_PanNET_PNET"),
    ("SPN", "04_SPN"),
    ("MCN", "05_MCN"),
    ("SCN", "06_SCN"),
    ("慢性胰腺炎", "07_慢性胰腺炎"),
    ("急性胰腺炎", "08_急性胰腺炎"),
)

RAW_FIELDS = (
    "姓名", "病理号", "审核医生", "肉眼所见", "标本类型", "标本名称", "送检科室", "送检医生",
    "临床诊断", "病理诊断", "病人类别", "病人编号", "申请序号", "住院号", "性别", "年龄",
    "婚姻", "民族", "身份证号", "联系信息", "收到日期", "取材日期", "镜下所见", "报告日期",
    "蜡块总数", "补充意见1", "补充意见2", "匹配日期",
)
HEADER = tuple(RAW_FIELDS)
CONTENT_FIELDS = ("肉眼所见", "病理诊断", "镜下所见", "补充意见1", "补充意见2")
ID_FIELD = "病人编号"
PATHOLOGY_NO_FIELD = "病理号"
ID_CARD_FIELD = "身份证号"
INPATIENT_FIELD = "住院号"
DATE_FIELDS = ("收到日期", "取材日期", "报告日期", "匹配日期")

EXPECTED = {
    "source_row_count": 9996,
    "pathology_record_count": 9528,
    "duplicate_source_row_count": 468,
    "linked_record_count": 9460,
    "identity_conflict_count": 62,
    "unmatched_count": 6,
    "time_sequence_conflict_count": 108,
    "received_missing_count": 244,
    "specimen_missing_count": 137,
    "report_missing_count": 244,
}

INT_FIELDS = {
    "source_row", "source_row_count", "identity_candidate_count", "pathology_record_count",
    "source_file_count", "source_row_count_total",
}
BOOL_FIELDS = {
    "pathology_time_sequence_conflict", "identity_eligible", "timeline_candidate_eligible",
    "event_eligible", "is_canonical", "is_content_conflict",
}


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and math.isnan(value):
        return ""
    return str(value)


def norm(value: Any, id_type: str = "PATIENT_ID") -> str | None:
    value = text(value).strip()
    if value.casefold() in MISSING:
        return None
    if id_type in {"PATIENT_ID", "VISIT_ID", "ID_CARD"}:
        return value.upper()
    return value


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def sha256_text(value: str) -> str:
    return sha256_bytes(value.encode("utf-8"))


def file_sha256(path: Path) -> tuple[int, str]:
    digest = hashlib.sha256()
    size = 0
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            size += len(block)
            digest.update(block)
    return size, digest.hexdigest()


def atomic_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", dir=path.parent, delete=False, suffix=".partial") as stream:
        json.dump(payload, stream, ensure_ascii=False, indent=2, sort_keys=True)
        stream.write("\n")
        temporary = Path(stream.name)
    os.replace(temporary, path)


def json_safe(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        return None if math.isnan(value) else value
    if isinstance(value, (datetime, date)):
        return {"type": type(value).__name__, "value": value.isoformat()}
    return {"type": type(value).__name__, "value": str(value)}


def cell_text(value: Any) -> str:
    return text(value)


def typed_values(values: Iterable[Any]) -> list[Any]:
    return [json_safe(value) for value in values]


def row_hash(values: Iterable[Any]) -> str:
    return sha256_text(json.dumps(typed_values(values), ensure_ascii=False, separators=(",", ":"), sort_keys=False))


def content_hash(record: dict[str, Any]) -> str:
    return row_hash(record.get(field) for field in CONTENT_FIELDS)


def source_record_key(file_sha: str, sheet_name: str, physical_row: int) -> str:
    return sha256_text(f"{file_sha}|{sheet_name}|{physical_row}")


def pathology_record_uid(pathology_no: str | None, source_key: str) -> str:
    identity = f"PATHOLOGY_NO|{pathology_no}" if pathology_no else f"SOURCE_RECORD|{source_key}"
    return str(uuid.uuid5(PATHOLOGY_NAMESPACE, identity))


def patient_uid(patient_id: str) -> str:
    return str(uuid.uuid5(PATIENT_NAMESPACE, f"PATIENT_ID|{patient_id}"))


def id_card_valid(value: Any) -> bool:
    value = norm(value, "ID_CARD")
    if value is None or not ID_CARD_RE.fullmatch(value):
        return False
    weights = (7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2)
    checks = "10X98765432"
    return checks[sum(int(a) * b for a, b in zip(value[:17], weights)) % 11] == value[-1]


def parse_date_value(value: Any) -> tuple[str, str]:
    raw = cell_text(value)
    if not raw:
        return "", "MISSING"
    if isinstance(value, datetime):
        return value.isoformat(sep=" "), "PARSED"
    if isinstance(value, date):
        return value.isoformat(), "PARSED"
    candidate = raw.strip().replace("/", "-").replace("年", "-").replace("月", "-").replace("日", "")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            parsed = datetime(1899, 12, 30) + timedelta(days=float(value))
            return parsed.isoformat(sep=" "), "PARSED"
        except (OverflowError, ValueError):
            return raw, "INVALID"
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S.%f"):
        try:
            return datetime.strptime(candidate[:26], fmt).isoformat(sep=" "), "PARSED"
        except ValueError:
            continue
    return raw, "INVALID"


def sequence_conflict(parsed: dict[str, str]) -> bool:
    ordered = [parsed.get("received_date_parsed", ""), parsed.get("specimen_date_parsed", ""), parsed.get("report_date_parsed", "")]
    values = []
    for value in ordered:
        if not value:
            continue
        try:
            values.append(datetime.fromisoformat(value))
        except ValueError:
            return True
    return any(left > right for left, right in zip(values, values[1:]))


def discover_inputs() -> list[dict[str, Any]]:
    inputs: list[dict[str, Any]] = []
    for disease_label, directory_name in DISEASE_INPUTS:
        directory = PATHOLOGY_ROOT / directory_name
        candidates = sorted(path for path in directory.glob("病理表*.xlsx") if path.is_file() and not path.name.startswith(".~lock"))
        if len(candidates) != 1:
            raise RuntimeError(f"PATHOLOGY_INPUT_COUNT:{directory_name}:{len(candidates)}")
        inputs.append({"disease_label": disease_label, "source_file": str(candidates[0].resolve()), "source_sheet": "Sheet1"})
    return inputs


def inspect_input(item: dict[str, Any]) -> dict[str, Any]:
    path = Path(item["source_file"])
    size, digest = file_sha256(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ["Sheet1"]:
            raise RuntimeError(f"PATHOLOGY_SHEET_SET:{path.name}")
        sheet = workbook["Sheet1"]
        iterator = sheet.iter_rows(values_only=True)
        actual_header = tuple(cell_text(value) for value in next(iterator, ()))
        if actual_header != HEADER:
            raise RuntimeError(f"PATHOLOGY_HEADER_MISMATCH:{path.name}")
        data_rows = 0
        for values in iterator:
            if len(values) != len(RAW_FIELDS):
                raise RuntimeError(f"PATHOLOGY_COLUMN_COUNT:{path.name}:{data_rows + 2}:{len(values)}")
            data_rows += 1
    finally:
        workbook.close()
    return {**item, "file_name": path.name, "file_size": size, "file_sha256": digest, "row_count": data_rows}


def excluded_files() -> dict[str, Any]:
    all_files = sorted(path for path in PATHOLOGY_ROOT.rglob("*") if path.is_file())
    allowed = {Path(item["source_file"]).resolve() for item in discover_inputs()}
    excluded = [path for path in all_files if path.resolve() not in allowed]
    lock_files = [path for path in excluded if path.name.startswith(".~lock") or ".~lock" in path.name]
    ct_mr_files = [path for path in excluded if path.suffix.lower() in {".xlsx", ".xlsm"} and ("CT" in path.name or "MR" in path.name)]
    return {
        "excluded_file_count": len(excluded),
        "lock_file_count": len(lock_files),
        "ct_mr_file_count": len(ct_mr_files),
        "excluded_files_by_name": sorted(path.name for path in excluded),
    }


def table_counts(db_path: Path) -> dict[str, int]:
    counts: dict[str, int] = {}
    with closing(sqlite3.connect(f"file:{db_path}?mode=ro", uri=True, timeout=30)) as db:
        statuses = db.execute("SELECT status,COUNT(*) FROM task_state GROUP BY status ORDER BY status").fetchall()
        counts["task_count"] = sum(int(row[1]) for row in statuses)
        counts.update({f"task_status_{row[0]}": int(row[1]) for row in statuses})
    return counts


def snapshot_stage6_base() -> dict[str, Any]:
    if not STAGE6_STATE.exists():
        raise RuntimeError("STAGE6_STATE_MISSING")
    state_size, state_sha = file_sha256(STAGE6_STATE)
    state_counts = table_counts(STAGE6_STATE)
    if state_counts.get("task_count") != 520 or state_counts.get("task_status_SUCCEEDED") != 520:
        raise RuntimeError(f"STAGE6_STATE_NOT_COMPLETE:{state_counts}")
    files = []
    if STAGE6_FULL_ROOT.exists():
        for path in sorted(path for path in STAGE6_FULL_ROOT.rglob("*") if path.is_file()):
            size, digest = file_sha256(path)
            files.append({"relative_path": str(path.relative_to(STAGE6_FULL_ROOT)), "size": size, "sha256": digest})
    return {"state_file": str(STAGE6_STATE), "state_size": state_size, "state_sha256": state_sha, "task_counts": state_counts, "formal_file_count": len(files), "formal_files": files}


def compare_stage6_base(snapshot: dict[str, Any]) -> dict[str, Any]:
    current = snapshot_stage6_base()
    return {"unchanged": current == snapshot, "expected": snapshot, "current": current}


def init_state(path: Path, manifest: dict[str, Any], phase: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with closing(sqlite3.connect(path)) as db:
        db.executescript(
            """
            CREATE TABLE IF NOT EXISTS metadata(key TEXT PRIMARY KEY,value TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS input_file(
                source_file TEXT PRIMARY KEY,disease_label TEXT NOT NULL,source_sheet TEXT NOT NULL,
                file_size INTEGER NOT NULL,file_sha256 TEXT NOT NULL,expected_row_count INTEGER NOT NULL,
                status TEXT NOT NULL,error_reason TEXT,updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS task_state(
                task_id TEXT PRIMARY KEY,phase TEXT NOT NULL,status TEXT NOT NULL,expected_row_count INTEGER NOT NULL,
                actual_row_count INTEGER,attempts INTEGER NOT NULL DEFAULT 0,error_reason TEXT,updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS output_file(
                output_kind TEXT NOT NULL,output_file TEXT PRIMARY KEY,row_count INTEGER NOT NULL,
                sha256 TEXT NOT NULL,schema_sha256 TEXT NOT NULL,status TEXT NOT NULL,updated_at TEXT NOT NULL
            );
            """
        )
        metadata = {
            "version": INGESTION_VERSION,
            "rule_version": RULE_VERSION,
            "phase": phase,
            "manifest_sha256": sha256_text(json.dumps(manifest, ensure_ascii=False, sort_keys=True)),
        }
        for key, value in metadata.items():
            db.execute("INSERT OR REPLACE INTO metadata(key,value) VALUES(?,?)", (key, value))
        for item in manifest["inputs"]:
            db.execute(
                "INSERT OR REPLACE INTO input_file(source_file,disease_label,source_sheet,file_size,file_sha256,expected_row_count,status,error_reason,updated_at) VALUES(?,?,?,?,?,?,?,?,datetime('now'))",
                (item["source_file"], item["disease_label"], item["source_sheet"], item["file_size"], item["file_sha256"], item["row_count"], "PENDING", None),
            )
        db.commit()


def read_source_rows(item: dict[str, Any]) -> list[dict[str, Any]]:
    path = Path(item["source_file"])
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows: list[dict[str, Any]] = []
    try:
        sheet = workbook[item["source_sheet"]]
        iterator = sheet.iter_rows(values_only=True)
        actual_header = tuple(cell_text(value) for value in next(iterator, ()))
        if actual_header != HEADER:
            raise RuntimeError(f"PATHOLOGY_HEADER_MISMATCH:{path.name}")
        for physical_row, values in enumerate(iterator, start=2):
            if len(values) != len(RAW_FIELDS):
                raise RuntimeError(f"PATHOLOGY_COLUMN_COUNT:{path.name}:{physical_row}:{len(values)}")
            raw = {field: values[index] for index, field in enumerate(RAW_FIELDS)}
            values_hash = row_hash(values)
            key = source_record_key(item["file_sha256"], item["source_sheet"], physical_row)
            rows.append({
                **{field: cell_text(raw[field]) for field in RAW_FIELDS},
                "raw_values_json": json.dumps(typed_values(values), ensure_ascii=False, separators=(",", ":")),
                "source_file": item["source_file"],
                "source_file_sha256": item["file_sha256"],
                "source_sheet": item["source_sheet"],
                "source_row": physical_row,
                "source_record_key": key,
                "row_hash": values_hash,
                "content_hash": content_hash(raw),
                "disease_label": item["disease_label"],
                "ingestion_version": INGESTION_VERSION,
            })
    finally:
        workbook.close()
    return rows


def build_model(manifest: dict[str, Any]) -> dict[str, Any]:
    source_rows: list[dict[str, Any]] = []
    for item in manifest["inputs"]:
        source_rows.extend(read_source_rows(item))
    groups: OrderedDict[str, dict[str, Any]] = OrderedDict()
    for row in source_rows:
        pathology_no = norm(row[PATHOLOGY_NO_FIELD], "VISIT_ID")
        group_key = f"PATHOLOGY_NO|{pathology_no}" if pathology_no else f"SOURCE_RECORD|{row['source_record_key']}"
        group = groups.setdefault(group_key, {"pathology_no": pathology_no, "rows": [], "diseases": set(), "row_hashes": set(), "content_hashes": set()})
        group["rows"].append(row)
        group["diseases"].add(row["disease_label"])
        group["row_hashes"].add(row["row_hash"])
        group["content_hashes"].add(row["content_hash"])
    records: list[dict[str, Any]] = []
    source_map: list[dict[str, Any]] = []
    labels: list[dict[str, Any]] = []
    for group in groups.values():
        canonical = group["rows"][0]
        uid = pathology_record_uid(group["pathology_no"], canonical["source_record_key"])
        for index, row in enumerate(group["rows"]):
            row["pathology_record_uid"] = uid
            source_map.append({
                "source_record_key": row["source_record_key"], "pathology_record_uid": uid,
                "pathology_no_normalized": group["pathology_no"] or "", "source_file": row["source_file"],
                "source_sheet": row["source_sheet"], "source_row": row["source_row"],
                "disease_label": row["disease_label"], "row_hash": row["row_hash"],
                "map_status": "CANONICAL" if index == 0 else ("CONTENT_CONFLICT" if len(group["content_hashes"]) > 1 else "DUPLICATE"),
                "ingestion_version": INGESTION_VERSION,
            })
            labels.append({
                "pathology_record_uid": uid, "source_record_key": row["source_record_key"],
                "disease_label": row["disease_label"], "label_source": "DISEASE_DIRECTORY",
                "ingestion_version": INGESTION_VERSION,
            })
        parsed: dict[str, str] = {}
        for field, output in (("收到日期", "received_date"), ("取材日期", "specimen_date"), ("报告日期", "report_date"), ("匹配日期", "match_date")):
            parsed_value, status = parse_date_value(canonical[field])
            parsed[f"{output}_parsed"] = parsed_value if status == "PARSED" else ""
            parsed[f"{output}_parse_status"] = status
        time_conflict = sequence_conflict(parsed)
        records.append({
            **canonical,
            "pathology_record_uid": uid,
            "pathology_no_normalized": group["pathology_no"] or "",
            "canonical_source_record_key": canonical["source_record_key"],
            "source_row_count": len(group["rows"]),
            "disease_labels_json": json.dumps(sorted(group["diseases"]), ensure_ascii=False),
            "is_content_conflict": len(group["content_hashes"]) > 1,
            "content_conflict_status": "PATHOLOGY_NO_CONTENT_CONFLICT" if len(group["content_hashes"]) > 1 else "NONE",
            "pathology_time_sequence_conflict": time_conflict,
            **parsed,
        })
    return {"source_rows": source_rows, "records": records, "source_map": source_map, "labels": labels}


def lookup_patient_uids(db: sqlite3.Connection, patient_id: str, cache: dict[str, list[str]]) -> list[str]:
    if patient_id in cache:
        return cache[patient_id]
    rows = db.execute("SELECT DISTINCT patient_uid FROM identity_alias WHERE global_identity_key=?", (f"PATIENT_ID|{patient_id}",)).fetchall()
    cache[patient_id] = sorted({row[0] for row in rows})
    return cache[patient_id]


def lookup_card_uids(db: sqlite3.Connection, card: str, cache: dict[str, dict[str, Any]]) -> dict[str, Any]:
    if card in cache:
        return cache[card]
    card_hash = sha256_text(card)
    blocked = db.execute("SELECT 1 FROM blocked_identity_card WHERE card_hash=? LIMIT 1", (card_hash,)).fetchone() is not None
    rows = db.execute("SELECT DISTINCT patient_uid FROM card_patient WHERE card_hash=?", (card_hash,)).fetchall()
    rows += db.execute("SELECT DISTINCT patient_uid FROM identity_alias WHERE id_type='ID_CARD' AND normalized_value=?", (card,)).fetchall()
    result = {"card_hash": card_hash, "blocked": blocked, "uids": sorted({row[0] for row in rows})}
    cache[card] = result
    return result


def resolve_records(model: dict[str, Any], stage6_state: Path = STAGE6_STATE) -> dict[str, Any]:
    links: list[dict[str, Any]] = []
    conflicts: list[dict[str, Any]] = []
    unmatched: list[dict[str, Any]] = []
    patient_delta: OrderedDict[str, dict[str, Any]] = OrderedDict()
    alias_delta: OrderedDict[str, dict[str, Any]] = OrderedDict()
    patient_cache: dict[str, list[str]] = {}
    card_cache: dict[str, dict[str, Any]] = {}
    encounter_cache: dict[tuple[str, str], str | None] = {}
    with closing(sqlite3.connect(f"file:{stage6_state}?mode=ro", uri=True, timeout=30)) as db:
        for record in model["records"]:
            patient_id = norm(record[ID_FIELD], "PATIENT_ID")
            card = norm(record[ID_CARD_FIELD], "ID_CARD")
            card = card if card and id_card_valid(card) else None
            patient_candidates = lookup_patient_uids(db, patient_id, patient_cache) if patient_id else []
            card_info = lookup_card_uids(db, card, card_cache) if card else {"card_hash": "", "blocked": False, "uids": []}
            card_candidates = card_info["uids"]
            # A multi-patient card is unsafe for card-based linking, but it is
            # not by itself a patient-number/card disagreement.  A trusted
            # patient number may still hard-link the record.  This keeps the
            # explicit conflict disposition limited to deterministic identity
            # disagreement, while ambiguous cards remain non-authoritative.
            conflict = bool(len(patient_candidates) > 1 or (len(patient_candidates) == 1 and len(card_candidates) == 1 and patient_candidates[0] != card_candidates[0]))
            patient = ""
            method = "PATIENT_UNMATCHED"
            disposition = "unmatched"
            link_status = "PATIENT_UNMATCHED"
            confidence = "NONE"
            if conflict:
                disposition = "conflict"
                link_status = "IDENTITY_CONFLICT_NO_AUTO_MERGE"
                method = "PATIENT_ID_CARD_DISAGREE_OR_BLOCKED"
            elif patient_candidates:
                patient = patient_candidates[0]
                disposition = "hard"
                link_status = "PATIENT_ID_EXACT"
                method = "PATIENT_ID_EXACT_TRUSTED_ALIAS"
                confidence = "HIGH"
            elif len(card_candidates) == 1 and not card_info["blocked"]:
                patient = card_candidates[0]
                disposition = "hard"
                link_status = "ID_CARD_RESCUE"
                method = "ID_CARD_EXACT_TRUSTED_ALIAS_NO_PATIENT_ALIAS_REGISTRATION"
                confidence = "HIGH"
            elif patient_id:
                patient = patient_uid(patient_id)
                disposition = "hard"
                link_status = "PATIENT_ID_NEW_DETERMINISTIC"
                method = "PATIENT_ID_DETERMINISTIC_NEW_INCREMENT_PATIENT"
                confidence = "HIGH"
                if patient_id not in patient_delta:
                    patient_delta[patient_id] = {"patient_uid": patient, "normalized_patient_id": patient_id, "source_record_key": record["canonical_source_record_key"], "rule_version": RULE_VERSION}
                    alias_key = f"pathology|PATIENT_ID|{patient_id}"
                    alias_delta[patient_id] = {"alias_key": alias_key, "global_identity_key": f"PATIENT_ID|{patient_id}", "source_system": "pathology", "id_type": "PATIENT_ID", "normalized_value": patient_id, "raw_value": record[ID_FIELD], "patient_uid": patient, "source_record_key": record["canonical_source_record_key"], "rule_version": RULE_VERSION}
            if patient and record[INPATIENT_FIELD]:
                visit = norm(record[INPATIENT_FIELD], "VISIT_ID")
                cache_key = (patient, visit or "")
                if cache_key not in encounter_cache:
                    row = db.execute("SELECT encounter_uid FROM encounter_registry WHERE patient_uid=? AND visit_id_normalized=? LIMIT 1", cache_key).fetchone()
                    encounter_cache[cache_key] = row[0] if row else None
                encounter = encounter_cache[cache_key] or ""
            else:
                encounter = ""
            identity_ok = disposition == "hard" and bool(patient)
            link = {
                "source_record_key": record["canonical_source_record_key"], "source_system": "pathology",
                "source_file": record["source_file"], "source_sheet": record["source_sheet"],
                "source_row": record["source_row"], "source_record_id": record["pathology_record_uid"],
                "pathology_record_uid": record["pathology_record_uid"], "patient_uid": patient,
                "encounter_uid": encounter, "link_status": link_status, "disposition_class": disposition,
                "link_method": method, "confidence": confidence, "identity_eligible": identity_ok,
                "timeline_candidate_eligible": identity_ok, "event_eligible": identity_ok,
                "pathology_time_sequence_conflict": record["pathology_time_sequence_conflict"],
                "encounter_match_status": "EXACT" if encounter else ("NOT_ATTEMPTED" if not record[INPATIENT_FIELD] else "NO_EXACT_MATCH"),
                "pathology_inpatient_no_evidence": record[INPATIENT_FIELD], "rule_version": RULE_VERSION,
            }
            links.append(link)
            if conflict:
                conflicts.append({
                    "pathology_record_uid": record["pathology_record_uid"], "source_record_key": record["canonical_source_record_key"],
                    "conflict_type": "IDENTITY_CARD_PATIENT_ID_CONFLICT", "resolution_status": "NO_AUTO_MERGE",
                    "identity_candidate_count": len(set(patient_candidates + card_candidates)),
                    "patient_candidate_hash": sha256_text("|".join(patient_candidates)),
                    "card_candidate_hash": sha256_text("|".join(card_candidates)), "rule_version": RULE_VERSION,
                })
            elif disposition == "unmatched":
                unmatched.append({
                    "pathology_record_uid": record["pathology_record_uid"], "source_record_key": record["canonical_source_record_key"],
                    "unmatched_reason": "NO_VALID_PATIENT_ID_OR_ID_CARD_MATCH", "pathology_inpatient_no_evidence": record[INPATIENT_FIELD],
                    "rule_version": RULE_VERSION,
                })
    return {"links": links, "conflicts": conflicts, "unmatched": unmatched, "patient_delta": list(patient_delta.values()), "alias_delta": list(alias_delta.values())}


def output_schema(fields: list[str]) -> pa.Schema:
    return pa.schema([(field, pa.int64() if field in INT_FIELDS else pa.bool_() if field in BOOL_FIELDS else pa.string()) for field in fields])


def write_parquet_rows(rows: list[dict[str, Any]], output_root: Path, kind: str, fields: list[str], *, max_rows: int = 50000) -> list[dict[str, Any]]:
    output_dir = output_root / kind
    output_dir.mkdir(parents=True, exist_ok=True)
    schema = output_schema(fields)
    results: list[dict[str, Any]] = []
    for start in range(0, len(rows), max_rows):
        chunk = rows[start:start + max_rows]
        shard = start // max_rows + 1
        final = output_dir / f"part-{shard:05d}.parquet"
        partial = output_dir / f"part-{shard:05d}.parquet.partial"
        if final.exists():
            raise FileExistsError(f"OUTPUT_EXISTS:{final}")
        arrays = [pa.array([row.get(field) for row in chunk], type=schema.field(field).type) for field in fields]
        table = pa.Table.from_arrays(arrays, schema=schema)
        pq.write_table(table, partial, compression="zstd")
        size, digest = file_sha256(partial)
        schema_sha = sha256_text(str(schema))
        os.replace(partial, final)
        results.append({"output_kind": kind, "output_file": str(final), "row_count": table.num_rows, "size": size, "sha256": digest, "schema_sha256": schema_sha, "status": "SUCCEEDED"})
    return results


SOURCE_ROW_FIELDS = list(RAW_FIELDS) + [
    "raw_values_json", "source_file", "source_file_sha256", "source_sheet", "source_row", "source_record_key", "row_hash", "content_hash", "disease_label", "pathology_record_uid", "ingestion_version",
]
RECORD_FIELDS = list(RAW_FIELDS) + [
    "raw_values_json", "source_file", "source_file_sha256", "source_sheet", "source_row", "source_record_key", "row_hash", "content_hash", "disease_label", "pathology_record_uid", "pathology_no_normalized", "canonical_source_record_key", "source_row_count", "disease_labels_json", "is_content_conflict", "content_conflict_status", "pathology_time_sequence_conflict", "received_date_parsed", "received_date_parse_status", "specimen_date_parsed", "specimen_date_parse_status", "report_date_parsed", "report_date_parse_status", "match_date_parsed", "match_date_parse_status", "ingestion_version",
]
SOURCE_MAP_FIELDS = ["source_record_key", "pathology_record_uid", "pathology_no_normalized", "source_file", "source_sheet", "source_row", "disease_label", "row_hash", "map_status", "ingestion_version"]
LABEL_FIELDS = ["pathology_record_uid", "source_record_key", "disease_label", "label_source", "ingestion_version"]
LINK_FIELDS = ["source_record_key", "source_system", "source_file", "source_sheet", "source_row", "source_record_id", "pathology_record_uid", "patient_uid", "encounter_uid", "link_status", "disposition_class", "link_method", "confidence", "identity_eligible", "timeline_candidate_eligible", "event_eligible", "pathology_time_sequence_conflict", "encounter_match_status", "pathology_inpatient_no_evidence", "rule_version"]
PATIENT_DELTA_FIELDS = ["patient_uid", "normalized_patient_id", "source_record_key", "rule_version"]
ALIAS_DELTA_FIELDS = ["alias_key", "global_identity_key", "source_system", "id_type", "normalized_value", "raw_value", "patient_uid", "source_record_key", "rule_version"]
CONFLICT_FIELDS = ["pathology_record_uid", "source_record_key", "conflict_type", "resolution_status", "identity_candidate_count", "patient_candidate_hash", "card_candidate_hash", "rule_version"]
UNMATCHED_FIELDS = ["pathology_record_uid", "source_record_key", "unmatched_reason", "pathology_inpatient_no_evidence", "rule_version"]


def select_canary_records(model: dict[str, Any], resolved: dict[str, Any]) -> list[dict[str, Any]]:
    required = {item["pathology_record_uid"] for item in resolved["links"] if item["disposition_class"] in {"conflict", "unmatched"}}
    required.update(item["pathology_record_uid"] for item in resolved["links"] if item["link_status"] == "ID_CARD_RESCUE")
    required.update(record["pathology_record_uid"] for record in model["records"] if record["source_row_count"] > 1 or len(json.loads(record["disease_labels_json"])) > 1)
    required.update(item["pathology_record_uid"] for item in resolved["links"] if item["link_status"] == "PATIENT_ID_NEW_DETERMINISTIC")
    selected = [record for record in model["records"] if record["pathology_record_uid"] in required]
    return sorted(selected, key=lambda row: row["pathology_record_uid"])


def audit_no_pii(payload: dict[str, Any]) -> bool:
    serialized = json.dumps(payload, ensure_ascii=False)
    forbidden = ("姓名", "身份证号", "联系信息", "病理诊断", "肉眼所见", "镜下所见", "补充意见1", "补充意见2")
    return not any(token in serialized for token in forbidden)


def build_manifest() -> dict[str, Any]:
    inputs = [inspect_input(item) for item in discover_inputs()]
    base = snapshot_stage6_base()
    excluded = excluded_files()
    manifest = {
        "manifest_version": "stage6_pathology_manifest_v1", "ingestion_version": INGESTION_VERSION, "rule_version": RULE_VERSION,
        "created_by": str(Path(__file__).resolve()), "inputs": inputs, "excluded": excluded,
        "expected": EXPECTED, "stage6_base_snapshot": base, "pii_in_manifest": False,
    }
    return manifest


def preflight() -> dict[str, Any]:
    if MANIFEST_PATH.exists() or INCREMENT_STATE.exists() or INCREMENT_FULL_ROOT.exists():
        raise RuntimeError("PATHOLOGY_INCREMENT_ALREADY_INITIALIZED")
    manifest = build_manifest()
    INCREMENT_ROOT.mkdir(parents=True, exist_ok=True)
    atomic_json(MANIFEST_PATH, manifest)
    init_state(INCREMENT_STATE, manifest, "full")
    report = {
        "report_version": "stage6_pathology_preflight_v1", "passed": True,
        "input_file_count": len(manifest["inputs"]), "source_row_count": sum(item["row_count"] for item in manifest["inputs"]),
        "input_files": [{"file_name": item["file_name"], "disease_label": item["disease_label"], "row_count": item["row_count"], "size": item["file_size"], "sha256": item["file_sha256"]} for item in manifest["inputs"]],
        "excluded": manifest["excluded"], "stage6_task_count": manifest["stage6_base_snapshot"]["task_counts"].get("task_count"),
        "stage6_succeeded_count": manifest["stage6_base_snapshot"]["task_counts"].get("task_status_SUCCEEDED"),
        "formal_file_count": manifest["stage6_base_snapshot"]["formal_file_count"], "pii_in_report": False,
    }
    if not audit_no_pii(report):
        raise RuntimeError("PII_IN_PREFLIGHT_REPORT")
    atomic_json(PREFLIGHT_REPORT, report)
    return report


def load_manifest() -> dict[str, Any]:
    if not MANIFEST_PATH.exists():
        raise RuntimeError("PATHOLOGY_MANIFEST_MISSING")
    return json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))


def verify_inputs_and_base(manifest: dict[str, Any]) -> dict[str, Any]:
    current_inputs = [inspect_input({"disease_label": item["disease_label"], "source_file": item["source_file"], "source_sheet": item["source_sheet"]}) for item in manifest["inputs"]]
    if current_inputs != manifest["inputs"]:
        raise RuntimeError("PATHOLOGY_INPUT_CHANGED")
    base_check = compare_stage6_base(manifest["stage6_base_snapshot"])
    if not base_check["unchanged"]:
        raise RuntimeError("STAGE6_BASE_CHANGED")
    return {"inputs_unchanged": True, "stage6_base_unchanged": True}


def run_canary() -> dict[str, Any]:
    manifest = load_manifest()
    verify_inputs_and_base(manifest)
    if CANARY_REPORT.exists() or INCREMENT_CANARY_ROOT.exists():
        raise RuntimeError("PATHOLOGY_CANARY_ALREADY_INITIALIZED")
    model = build_model(manifest)
    resolved = resolve_records(model)
    selected = select_canary_records(model, resolved)
    selected_uids = {record["pathology_record_uid"] for record in selected}
    selected_links = [link for link in resolved["links"] if link["pathology_record_uid"] in selected_uids]
    canary = {"source_rows": [row for row in model["source_rows"] if row["pathology_record_uid"] in selected_uids], "records": selected, "source_map": [row for row in model["source_map"] if row["pathology_record_uid"] in selected_uids], "labels": [row for row in model["labels"] if row["pathology_record_uid"] in selected_uids]}
    resolved_canary = {**resolved, "links": selected_links, "conflicts": [row for row in resolved["conflicts"] if row["pathology_record_uid"] in selected_uids], "unmatched": [row for row in resolved["unmatched"] if row["pathology_record_uid"] in selected_uids]}
    canary_outputs: list[dict[str, Any]] = []
    for rows, kind, fields in ((canary["source_rows"], "pathology_source_row", SOURCE_ROW_FIELDS), (canary["records"], "pathology_record", RECORD_FIELDS), (canary["source_map"], "pathology_source_map", SOURCE_MAP_FIELDS), (canary["labels"], "pathology_disease_label", LABEL_FIELDS), (resolved_canary["links"], "record_links_pathology", LINK_FIELDS), (resolved_canary["patient_delta"], "patient_master_delta", PATIENT_DELTA_FIELDS), (resolved_canary["alias_delta"], "identity_alias_delta", ALIAS_DELTA_FIELDS), (resolved_canary["conflicts"], "pathology_conflict_log", CONFLICT_FIELDS), (resolved_canary["unmatched"], "pathology_unmatched", UNMATCHED_FIELDS)):
        canary_outputs.extend(write_parquet_rows(rows, INCREMENT_CANARY_ROOT, kind, fields))
    counts = {
        "selected_source_row_count": len(canary["source_rows"]), "selected_pathology_record_count": len(canary["records"]),
        "selected_link_count": len(resolved_canary["links"]), "selected_conflict_count": len(resolved_canary["conflicts"]),
        "selected_unmatched_count": len(resolved_canary["unmatched"]), "selected_new_patient_count": len(resolved_canary["patient_delta"]),
        "selected_id_card_rescue_count": sum(link["link_status"] == "ID_CARD_RESCUE" for link in selected_links),
        "selected_duplicate_record_count": sum(record["source_row_count"] > 1 for record in selected),
    }
    passed = all(counts[key] > 0 for key in ("selected_conflict_count", "selected_unmatched_count", "selected_new_patient_count", "selected_id_card_rescue_count", "selected_duplicate_record_count"))
    report = {"report_version": "stage6_pathology_canary_v1", "canary_passed": passed, **counts, "output_count": len(canary_outputs), "pii_in_report": False}
    if not audit_no_pii(report):
        raise RuntimeError("PII_IN_CANARY_REPORT")
    atomic_json(CANARY_REPORT, report)
    return report


def _record_task_state(path: Path, task_id: str, status: str, expected: int, actual: int | None = None, error: str | None = None) -> None:
    with closing(sqlite3.connect(path)) as db:
        db.execute("INSERT OR REPLACE INTO task_state(task_id,phase,status,expected_row_count,actual_row_count,attempts,error_reason,updated_at) VALUES(?,?,?,?,?,?,?,datetime('now'))", (task_id, "full", status, expected, actual, 1, error))
        db.commit()


def run_full() -> dict[str, Any]:
    manifest = load_manifest()
    verify_inputs_and_base(manifest)
    if not CANARY_REPORT.exists() or not json.loads(CANARY_REPORT.read_text(encoding="utf-8")).get("canary_passed"):
        raise RuntimeError("PATHOLOGY_CANARY_NOT_PASSED")
    if INCREMENT_FULL_ROOT.exists():
        raise RuntimeError("PATHOLOGY_FULL_OUTPUT_ALREADY_EXISTS")
    model = build_model(manifest)
    _record_task_state(INCREMENT_STATE, "pathology:source_rows", "RUNNING", len(model["source_rows"]))
    resolved = resolve_records(model)
    output_specs = (
        (model["source_rows"], "pathology_source_row", SOURCE_ROW_FIELDS), (model["records"], "pathology_record", RECORD_FIELDS),
        (model["source_map"], "pathology_source_map", SOURCE_MAP_FIELDS), (model["labels"], "pathology_disease_label", LABEL_FIELDS),
        (resolved["links"], "record_links_pathology", LINK_FIELDS), (resolved["patient_delta"], "patient_master_delta", PATIENT_DELTA_FIELDS),
        (resolved["alias_delta"], "identity_alias_delta", ALIAS_DELTA_FIELDS), (resolved["conflicts"], "pathology_conflict_log", CONFLICT_FIELDS),
        (resolved["unmatched"], "pathology_unmatched", UNMATCHED_FIELDS),
    )
    outputs: list[dict[str, Any]] = []
    try:
        for rows, kind, fields in output_specs:
            outputs.extend(write_parquet_rows(rows, INCREMENT_FULL_ROOT, kind, fields))
    except Exception as exc:
        _record_task_state(INCREMENT_STATE, "pathology:source_rows", "FAILED", len(model["source_rows"]), sum(item["row_count"] for item in outputs), f"{type(exc).__name__}:{str(exc)[:160]}")
        raise
    _record_task_state(INCREMENT_STATE, "pathology:source_rows", "SUCCEEDED", len(model["source_rows"]), len(model["source_rows"]))
    _record_task_state(INCREMENT_STATE, "pathology:records", "SUCCEEDED", len(model["records"]), len(model["records"]))
    _record_task_state(INCREMENT_STATE, "pathology:record_links", "SUCCEEDED", len(model["records"]), len(resolved["links"]))
    with closing(sqlite3.connect(INCREMENT_STATE)) as db:
        for output in outputs:
            db.execute("INSERT OR REPLACE INTO output_file(output_kind,output_file,row_count,sha256,schema_sha256,status,updated_at) VALUES(?,?,?,?,?,?,datetime('now'))", (output["output_kind"], output["output_file"], output["row_count"], output["sha256"], output["schema_sha256"], "SUCCEEDED"))
        db.commit()
    input_count = len(model["source_rows"])
    report = {
        "report_version": "stage6_pathology_increment_v1", "full_run_completed": True,
        "source_row_count": input_count, "pathology_record_count": len(model["records"]),
        "duplicate_source_row_count": input_count - len(model["records"]), "linked_record_count": sum(link["disposition_class"] == "hard" for link in resolved["links"]),
        "identity_conflict_count": len(resolved["conflicts"]), "unmatched_count": len(resolved["unmatched"]),
        "id_card_rescue_count": sum(link["link_status"] == "ID_CARD_RESCUE" for link in resolved["links"]),
        "new_patient_count": len(resolved["patient_delta"]), "encounter_exact_match_count": sum(bool(link["encounter_uid"]) for link in resolved["links"]),
        "time_sequence_conflict_count": sum(record["pathology_time_sequence_conflict"] * record["source_row_count"] for record in model["records"]),
        "received_missing_count": sum(record["received_date_parse_status"] == "MISSING" for record in model["records"] for _ in range(record["source_row_count"])),
        "specimen_missing_count": sum(record["specimen_date_parse_status"] == "MISSING" for record in model["records"] for _ in range(record["source_row_count"])),
        "report_missing_count": sum(record["report_date_parse_status"] == "MISSING" for record in model["records"] for _ in range(record["source_row_count"])),
        "content_conflict_count": sum(record["is_content_conflict"] for record in model["records"]),
        "outputs": outputs, "stage6_base_unchanged": compare_stage6_base(manifest["stage6_base_snapshot"])["unchanged"],
        "state_database": str(INCREMENT_STATE), "pii_in_report": False,
    }
    if not audit_no_pii(report):
        raise RuntimeError("PII_IN_FULL_REPORT")
    atomic_json(FULL_REPORT, report)
    return report


def validate() -> dict[str, Any]:
    manifest = load_manifest()
    if not FULL_REPORT.exists():
        raise RuntimeError("PATHOLOGY_FULL_REPORT_MISSING")
    report = json.loads(FULL_REPORT.read_text(encoding="utf-8"))
    base_check = compare_stage6_base(manifest["stage6_base_snapshot"])
    output_files = []
    if INCREMENT_FULL_ROOT.exists():
        output_files = sorted(path for path in INCREMENT_FULL_ROOT.rglob("*.parquet") if path.is_file())
    file_checks = []
    for path in output_files:
        metadata = pq.read_metadata(path)
        size, digest = file_sha256(path)
        file_checks.append({"file": str(path), "row_count": int(metadata.num_rows), "size": size, "sha256": digest})
    expected_match = all(report.get(key) == value for key, value in EXPECTED.items() if key in report)
    source_conservation = report.get("source_row_count") == report.get("pathology_record_count", 0) + report.get("duplicate_source_row_count", 0)
    passed = bool(report.get("full_run_completed") and expected_match and source_conservation and base_check["unchanged"] and not any(path.name.endswith(".partial") for path in INCREMENT_FULL_ROOT.rglob("*")) and audit_no_pii(report))
    result = {"report_version": "stage6_pathology_validation_v1", "validation_passed": passed, "expected_match": expected_match, "source_conservation": source_conservation, "stage6_base_unchanged": base_check["unchanged"], "output_file_count": len(file_checks), "output_files": file_checks, "pii_in_report": False}
    if not audit_no_pii(result):
        raise RuntimeError("PII_IN_VALIDATION_REPORT")
    atomic_json(VALIDATION_REPORT, result)
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Stage 6 pathology increment v1")
    parser.add_argument("command", choices=("preflight", "canary", "full-run", "validate"))
    args = parser.parse_args(argv)
    if args.command == "preflight":
        result = preflight()
    elif args.command == "canary":
        result = run_canary()
    elif args.command == "full-run":
        result = run_full()
    else:
        result = validate()
    print(json.dumps({key: value for key, value in result.items() if key not in {"outputs", "output_files"}}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
