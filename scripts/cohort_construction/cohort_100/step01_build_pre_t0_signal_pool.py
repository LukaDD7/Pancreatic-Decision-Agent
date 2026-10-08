import csv
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from scripts.cohort_construction.paths import data_root

ROOT = data_root()
STAGE8 = ROOT / "pipeline_outputs_stage8_v1"
INDEX = STAGE8 / "pancreas_imaging_report_index_v2" / "restricted" / "index" / "report_index.parquet"
DOCUMENTS = ROOT / "患者入院出院文书.csv"
PATHOLOGY = (
    ROOT
    / "pipeline_outputs_stage6_v2"
    / "restricted"
    / "increment"
    / "pathology_total_v2"
    / "full"
    / "pathology_record"
    / "part-00001.parquet"
)
OUT = STAGE8 / "pancreas_decision_window_enriched_v1" / "restricted"
OUT.mkdir(parents=True, exist_ok=True)


STRONG = re.compile(
    r"转移\s*(?:待排|可能|不排除|[?？])|不除外\s*转移|可疑\s*转移|"
    r"考虑\s*(?:为)?\s*转移|待排\s*转移"
)
NEGATED = re.compile(r"未见|排除|不考虑")
BENIGN = re.compile(r"血管瘤|脉管瘤|囊肿|钙化|结石|脂肪浸润|反应性增生")
POST_TREATMENT = re.compile(r"术后|治疗后|化疗后|放疗后|复发")
WEAK = re.compile(r"结节|低强化灶|弥散受限|占位|异常灶")
WEAK_ACTION = re.compile(r"随访|复查|结合临床|建议.{0,16}(?:MR|MRI|磁共振|PET|增强)")
PANCREAS_METHOD = re.compile(r"胰腺.*(?:CT|MR)|胰胆管MR|MRCP|胰腺动脉CT", re.I)
PANCREAS_PATH = re.compile(r"胰(?:头|体|尾|腺)[^，,；;。\n]{0,24}(?:占位|肿瘤|癌|恶性|导管腺癌)")
PANCREAS_DOC_PATH = re.compile(
    r"胰(?:头|体|尾|腺)[^。；\n]{0,30}(?:占位|肿瘤|癌|恶性)|梗阻性黄疸|"
    r"(?:EUS|超声内镜|内镜超声|FNA|细针穿刺)|胰腺MDT|胰腺多学科"
)
DISTANT_SITE = re.compile(r"肝(?!门)|腹膜(?!后)|网膜|肺|胸膜|骨|锁骨上|纵隔|主动脉旁|肾上腺")
REGIONAL_NODE_ONLY = re.compile(r"肝门|腹腔干|肠系膜上动脉|SMA|腹膜后[^，,；;。\n]{0,10}淋巴结")
EXPLICIT_M1 = re.compile(
    r"(?:肝(?!门)|腹膜(?!后)|网膜|肺|胸膜|骨|锁骨上|纵隔|主动脉旁|肾上腺)"
    r"[^，,；;。\n]{0,20}(?:转移瘤|转移灶|转移癌|多发转移)"
)
SURGERY = re.compile(
    r"胰十二指肠切除|胰体尾(?:脾脏)?切除|远端胰腺切除|胰腺次全切除|全胰切除|"
    r"胰腺根治性大部切除|胰体癌根治|胰尾癌根治|剖腹探查|开腹探查"
)
RADICAL = re.compile(
    r"胰十二指肠切除|胰体尾(?:脾脏)?切除|远端胰腺切除|胰腺次全切除|全胰切除|"
    r"胰腺根治性大部切除|胰体癌根治|胰尾癌根治"
)
EXPLORE = re.compile(r"探查|活检|结节切除")
ACTUAL_SURGERY_RECORD = re.compile(r"实施手术|手术经过|术中诊断|术后首次病程记录")
SURGERY_DATE = re.compile(r"手术时间\s*[:：]?\s*(20\d{2})[-/]([01]?\d)[-/]([0-3]?\d)")
CHEMO = re.compile(
    r"今日开始给予[^。；\n]{0,100}(?:方案)?化疗|"
    r"入院后于\s*20\d{2}[-年/]\d{1,2}[-月/]\d{1,2}日?[^。；\n]{0,100}(?:方案)?(?:化疗|治疗)|"
    r"本次入院[^。；\n]{0,80}(?:方案)?第\s*\d+\s*次治疗|"
    r"诊疗经过[^。；\n]{0,120}(?:FOLFIRINOX|NALIRIFOX|FOLFIRI|AG方案|GS方案|吉西他滨|白蛋白紫杉醇)[^。；\n]{0,50}(?:化疗|治疗)",
    re.I,
)
MDT = re.compile(r"MDT|多学科门诊|多学科会诊", re.I)
DECISION_CANCEL = re.compile(r"取消手术|暂缓手术|放弃手术")
NO_SURGERY = re.compile(r"无手术指征|暂无手术指征|不宜手术")
DECISION_SYSTEMIC = re.compile(r"改行化疗|建议先行化疗|先行系统治疗|转化治疗")


def compact(value, limit=360):
    return re.sub(r"\s+", " ", str(value or "")).strip()[:limit]


def site_name(text):
    mapping = [
        (r"肝(?!门)", "肝"),
        (r"腹膜(?!后)|网膜", "腹膜/网膜"),
        (r"肺|胸膜", "肺/胸膜"),
        (r"骨", "骨"),
        (r"肾上腺", "肾上腺"),
        (r"锁骨上|纵隔|主动脉旁", "远处淋巴结"),
    ]
    for pattern, value in mapping:
        if re.search(pattern, text):
            return value
    return "其他"


def explicit_m1_in_conclusion(text):
    for clause in re.split(r"[，,；;。\n]+", str(text or "")):
        if EXPLICIT_M1.search(clause) and not STRONG.search(clause):
            return True
    return False


def extract_signal(text):
    for line in re.split(r"[\r\n；;。]+", str(text or "")):
        for match in STRONG.finditer(line):
            before = line[max(0, match.start() - 15) : match.start()]
            # The distant site must occur before the uncertainty phrase.  This
            # prevents a regional-node phrase such as "淋巴结转移可能，肝内胆管"
            # from borrowing an unrelated site that appears later in the line.
            before_site = line[max(0, match.start() - 45) : match.start()]
            local = line[max(0, match.start() - 45) : min(len(line), match.end() + 20)]
            if NEGATED.search(before):
                return None, "前15字含否定"
            if BENIGN.search(local):
                return None, "局部语境命中明确良性词"
            if not DISTANT_SITE.search(before_site):
                if REGIONAL_NODE_ONLY.search(local):
                    return None, "仅区域淋巴结/N1"
                return None, "未定位到高纯度远处部位"
            return compact(line, 600), ""
    return None, "未命中强信号"


def report_citation(row, quote):
    return f"{row['index_report_uid']}@{str(row['_exam_dt'])[:19]}:{compact(quote, 300)}"


def parse_date(value):
    try:
        return pd.Timestamp(value)
    except Exception:
        return pd.NaT


def clip_around(text, match, before=100, after=260):
    return compact(text[max(0, match.start() - before) : match.end() + after], before + after + 100)


def disease_from_pathology(text):
    text = str(text or "")
    if re.search(r"导管腺癌|PDAC", text, re.I):
        return "PDAC"
    if re.search(r"神经内分泌|PanNET|NET", text, re.I):
        return "PanNET"
    if re.search(r"IPMN|导管内乳头状黏液", text, re.I):
        return "IPMN"
    if re.search(r"慢性胰腺炎|良性|炎性", text):
        return "非肿瘤/良性"
    return "其他/待定"


def site_specific_negative_pathology(signal_site, pathology_text):
    site_patterns = {
        "肝": r"肝",
        "腹膜/网膜": r"腹膜|网膜",
        "肺/胸膜": r"肺|胸膜",
        "骨": r"骨",
        "肾上腺": r"肾上腺",
        "远处淋巴结": r"锁骨上|纵隔|主动脉旁",
    }
    site_pattern = site_patterns.get(signal_site)
    if not site_pattern:
        return False
    return bool(
        re.search(
            rf"(?:{site_pattern})[^，,；;。\n]{{0,12}}(?:未见(?:癌|肿瘤|恶性)|[（(][-－][)）])",
            pathology_text,
        )
    )


def style_workbook(writer):
    for worksheet in writer.book.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for index, column in enumerate(worksheet.columns, 1):
            values = [str(cell.value or "") for cell in list(column)[:200]]
            width = min(max(max((len(value) for value in values), default=8) + 2, 10), 48)
            worksheet.column_dimensions[get_column_letter(index)].width = width


def main():
    image = pd.read_parquet(INDEX)
    image["patient_id"] = image["patient_id"].astype(str)
    image["_exam_dt"] = pd.to_datetime(image["exam_datetime"], errors="coerce")
    image = image.sort_values(["patient_id", "_exam_dt"], na_position="last")

    discard = Counter()
    weak_rows = []
    report_signal = {}
    for index, row in image.iterrows():
        conclusion = str(row.get("description_deidentified") or "")
        if pd.isna(row["_exam_dt"]):
            if STRONG.search(conclusion):
                discard["报告日期/时间缺失"] += 1
            continue
        signal, reason = extract_signal(conclusion)
        if signal:
            if POST_TREATMENT.search(conclusion):
                discard["术后/治疗后信号"] += 1
                continue
            if explicit_m1_in_conclusion(conclusion):
                discard["同份结论已存在明确M1"] += 1
                continue
            report_signal[index] = signal
        elif STRONG.search(conclusion):
            discard[reason] += 1
        elif WEAK.search(conclusion) and WEAK_ACTION.search(conclusion):
            weak_rows.append(
                {
                    "患者ID": row["patient_id"],
                    "检查时间_非T0": str(row["_exam_dt"]),
                    "检查方法": row.get("exam_method", ""),
                    "报告ID": row.get("index_report_uid", ""),
                    "结论原文": compact(conclusion, 800),
                    "入weak_pool原因": "结论有弱病灶/复查建议，但无规定的强转移措辞",
                }
            )

    patient_signals = {}
    denominator_evidence = {}
    for patient_id, group in image.groupby("patient_id", sort=False):
        prior_pancreas = []
        prior_explicit_m1 = False
        patient_hit = False
        for index, row in group.iterrows():
            conclusion = str(row.get("description_deidentified") or "")
            if index in report_signal and not patient_hit:
                if prior_explicit_m1:
                    discard["首次强信号前已有明确M1"] += 1
                elif not prior_pancreas:
                    discard["首次强信号前无独立胰腺外科评估路径证据"] += 1
                else:
                    quote = report_signal[index]
                    patient_signals[patient_id] = {
                        "patient_id": patient_id,
                        "signal_exam_datetime": row["_exam_dt"],
                        "signal_date": str(row["_exam_dt"].date()),
                        "signal_site": site_name(quote),
                        "signal_quote": quote,
                        "signal_report_id": row.get("index_report_uid", ""),
                        "signal_exam_method": row.get("exam_method", ""),
                        "signal_citation": report_citation(row, quote),
                        "t0_valid": False,
                        "t0_blocker": "仅有检查时间，缺少报告签发/审核时间",
                    }
                    denominator_evidence[patient_id] = prior_pancreas[-1]
                    patient_hit = True
            if PANCREAS_METHOD.search(str(row.get("exam_method") or "")) or PANCREAS_PATH.search(conclusion):
                prior_pancreas.append(
                    {
                        "date": str(row["_exam_dt"]),
                        "type": "首次信号前胰腺协议影像/胰腺病变影像",
                        "citation": report_citation(row, conclusion),
                    }
                )
            if explicit_m1_in_conclusion(conclusion):
                prior_explicit_m1 = True

    # Downstream documents are scanned only for patients that passed the
    # pre-signal denominator, preserving the entry criterion independently of
    # whether surgery was eventually performed.
    document_events = defaultdict(list)
    with DOCUMENTS.open("r", encoding="gb18030", errors="replace", newline="") as handle:
        reader = csv.DictReader(handle)
        for row_number, row in enumerate(reader, start=2):
            patient_id = str(row.get("PATIENT_ID") or "")
            if patient_id not in patient_signals:
                continue
            create_time = parse_date(row.get("CREATE_DATE_TIME"))
            if pd.isna(create_time):
                continue
            signal_time = patient_signals[patient_id]["signal_exam_datetime"]
            days = (create_time.normalize() - signal_time.normalize()).days
            if days < -180 or days > 365:
                continue
            name = str(row.get("文书名称") or "")
            text = str(row.get("文书内容") or "")
            source = f"患者入院出院文书.csv:{row_number}@{create_time.date()}"
            if create_time < signal_time and PANCREAS_DOC_PATH.search(text):
                match = PANCREAS_DOC_PATH.search(text)
                document_events[patient_id].append(
                    {
                        "kind": "denominator_document",
                        "date": create_time,
                        "citation": f"{source}:{clip_around(text, match)}",
                    }
                )
            if 0 <= days <= 90:
                if ACTUAL_SURGERY_RECORD.search(name + " " + text) and SURGERY.search(text):
                    operation_match = SURGERY_DATE.search(text)
                    if operation_match:
                        operation_date = pd.Timestamp(
                            year=int(operation_match.group(1)),
                            month=int(operation_match.group(2)),
                            day=int(operation_match.group(3)),
                        )
                    else:
                        operation_date = create_time.normalize()
                    operation_days = (operation_date - signal_time.normalize()).days
                    if 1 <= operation_days <= 90:
                        procedure = SURGERY.search(text)
                        document_events[patient_id].append(
                            {
                                "kind": "operation",
                                "date": operation_date,
                                "days": operation_days,
                                "radical": bool(RADICAL.search(text)),
                                "exploration": bool(EXPLORE.search(text)),
                                "citation": f"{source}:{clip_around(text, procedure)}",
                            }
                        )
                for kind, pattern in (
                    ("chemotherapy", CHEMO),
                    ("mdt", MDT),
                    ("cancel", DECISION_CANCEL),
                    ("no_surgery", NO_SURGERY),
                    ("systemic_decision", DECISION_SYSTEMIC),
                ):
                    match = pattern.search(text)
                    if match:
                        document_events[patient_id].append(
                            {
                                "kind": kind,
                                "date": create_time,
                                "days": days,
                                "citation": f"{source}:{clip_around(text, match)}",
                            }
                        )

    pathology = pd.read_parquet(
        PATHOLOGY,
        columns=["病人编号", "报告日期", "病理诊断", "标本名称", "source_record_key"],
    )
    pathology["病人编号"] = pathology["病人编号"].astype(str)
    pathology["_date"] = pd.to_datetime(pathology["报告日期"], errors="coerce")

    image_by_patient = {patient_id: group for patient_id, group in image.groupby("patient_id", sort=False)}
    candidates = []
    for patient_id, signal in patient_signals.items():
        events = document_events.get(patient_id, [])
        operations = sorted((event for event in events if event["kind"] == "operation"), key=lambda event: event["date"])
        cancellations = sorted((event for event in events if event["kind"] == "cancel"), key=lambda event: event["date"])
        systemic = sorted(
            (event for event in events if event["kind"] in {"systemic_decision", "chemotherapy"}),
            key=lambda event: event["date"],
        )
        no_surgery = sorted((event for event in events if event["kind"] == "no_surgery"), key=lambda event: event["date"])
        decision_event_date = (
            operations[0]["date"]
            if operations
            else cancellations[0]["date"]
            if cancellations
            else systemic[0]["date"]
            if systemic
            else signal["signal_exam_datetime"] + pd.Timedelta(days=90)
        )

        added_reports = []
        for _, report in image_by_patient[patient_id].iterrows():
            report_time = report["_exam_dt"]
            if not (signal["signal_exam_datetime"] < report_time <= decision_event_date):
                continue
            method = str(report.get("exam_method") or "")
            if re.search(r"肝脏MR|PET", method, re.I):
                quote = str(report.get("description_deidentified") or "")
                added_reports.append(report_citation(report, quote))

        mdt_events = sorted((event for event in events if event["kind"] == "mdt"), key=lambda event: event["date"])
        added_evaluation = "无"
        evaluation_evidence = []
        if added_reports:
            added_evaluation = "有"
            evaluation_evidence.extend(added_reports)
        if mdt_events:
            added_evaluation = "有"
            evaluation_evidence.extend(event["citation"] for event in mdt_events[:2])

        if cancellations:
            decision_change = "取消"
            decision_evidence = cancellations[0]["citation"]
        elif systemic and no_surgery and (not operations or systemic[0]["date"] < operations[0]["date"]):
            decision_change = "改系统治疗"
            decision_evidence = no_surgery[0]["citation"] + " | " + systemic[0]["citation"]
        elif operations and operations[0]["exploration"] and not operations[0]["radical"]:
            decision_change = "改探查活检"
            decision_evidence = operations[0]["citation"]
        elif operations and operations[0]["radical"]:
            decision_change = "继续根治切除"
            decision_evidence = operations[0]["citation"]
        else:
            decision_change = "无法判断"
            decision_evidence = ""

        p = pathology[pathology["病人编号"].eq(patient_id)].copy()
        p = p[
            (p["_date"] >= signal["signal_exam_datetime"].normalize())
            & (p["_date"] <= signal["signal_exam_datetime"].normalize() + pd.Timedelta(days=180))
        ].sort_values("_date")
        pathology_text = " ".join((p["标本名称"].fillna("") + " " + p["病理诊断"].fillna("")).tolist())
        disease = disease_from_pathology(pathology_text)
        m_state = "无法判断"
        m_evidence = ""
        post_signal_reports = image_by_patient[patient_id]
        post_signal_reports = post_signal_reports[
            post_signal_reports["_exam_dt"] > signal["signal_exam_datetime"]
        ].sort_values("_exam_dt")
        explicit_m1_reports = post_signal_reports[
            post_signal_reports["description_deidentified"].fillna("").map(explicit_m1_in_conclusion)
        ]
        if re.search(
            r"M\s*[:：]?\s*1|(?:肝|腹膜|网膜|肺|胸膜|骨|肾上腺)[^。；\n]{0,30}转移|"
            r"(?:腹膜|网膜|系膜)[^。；\n]{0,20}结节[^。；\n]{0,20}(?:癌组织|腺癌)",
            pathology_text,
            re.I,
        ):
            m_state = "M1"
            first = p.iloc[0]
            m_evidence = (
                f"{first['source_record_key']}@{first['_date'].date()}:"
                f"{compact(str(first['标本名称']) + ' ' + str(first['病理诊断']), 420)}"
            )
        elif not explicit_m1_reports.empty:
            first = explicit_m1_reports.iloc[0]
            m_state = "M1"
            m_evidence = report_citation(first, str(first.get("description_deidentified") or ""))
        elif site_specific_negative_pathology(signal["signal_site"], pathology_text):
            m_state = "M0"
            for _, pathology_row in p.iterrows():
                row_text = str(pathology_row["标本名称"]) + " " + str(pathology_row["病理诊断"])
                if site_specific_negative_pathology(signal["signal_site"], row_text):
                    m_evidence = (
                        f"{pathology_row['source_record_key']}@{pathology_row['_date'].date()}:"
                        f"{compact(row_text, 420)}"
                    )
                    break
        else:
            # A follow-up-based M0 requires an observation at least six months
            # after the signal and an explicit statement that distant
            # metastasis was not seen.  Merely having later imaging is not
            # enough.
            followup = image_by_patient[patient_id]
            followup = followup[
                (followup["_exam_dt"] >= signal["signal_exam_datetime"] + pd.Timedelta(days=180))
                & (followup["_exam_dt"] <= signal["signal_exam_datetime"] + pd.Timedelta(days=365))
            ].sort_values("_exam_dt")
            for _, report in followup.iterrows():
                conclusion = str(report.get("description_deidentified") or "")
                if (
                    re.search(r"远处转移\s*[:：]\s*未见|未见[^。；\n]{0,25}远处转移", conclusion)
                    and not explicit_m1_in_conclusion(conclusion)
                ):
                    m_state = "M0"
                    m_evidence = report_citation(report, conclusion)
                    break

        # The development set uses a second, explicitly weaker endpoint tier.
        # A patient with staging-capable imaging at least six months later and
        # no explicit M1 is useful for decision modelling, but is not relabelled
        # as M0.  This keeps the strict validation endpoint intact while avoiding
        # an unusably small development set.
        followup_observation = post_signal_reports[
            (post_signal_reports["_exam_dt"] >= signal["signal_exam_datetime"] + pd.Timedelta(days=180))
            & post_signal_reports["exam_method"].fillna("").str.contains(r"CT|MR|MRI|PET|磁共振", case=False, regex=True)
            & post_signal_reports["description_deidentified"].fillna("").ne("")
        ]
        if m_state == "M1":
            development_endpoint = "明确M1"
            development_evidence = m_evidence
        elif m_state == "M0":
            development_endpoint = "明确M0"
            development_evidence = m_evidence
        elif not followup_observation.empty:
            first_followup = followup_observation.iloc[0]
            development_endpoint = "≥6月未观察到明确M1（代理阴性，非M0）"
            development_evidence = report_citation(
                first_followup, str(first_followup.get("description_deidentified") or "")
            )
        else:
            development_endpoint = "结局待定"
            development_evidence = ""

        if m_state == "M1":
            information_gain = "确认 M1" if added_evaluation == "有" else "未追加评估"
        elif m_state == "M0":
            information_gain = "排除 M1" if added_evaluation == "有" else "未追加评估"
        else:
            information_gain = "仍不确定" if added_evaluation == "有" else "未追加评估"

        if added_evaluation == "无" and m_state == "M1":
            proxy = "术中意外 M1 代理阳性"
        else:
            proxy = "不适用"

        content_confidence = "高" if decision_change != "无法判断" and m_state != "无法判断" else "中" if decision_change != "无法判断" else "低"
        final_confidence = "低"
        candidates.append(
            {
                "患者ID": patient_id,
                "首次信号日期_检查日期暂代": signal["signal_date"],
                "正式T0": "缺失",
                "T0状态": signal["t0_blocker"],
                "部位": signal["signal_site"],
                "原文": signal["signal_quote"],
                "信号证据引用": signal["signal_citation"],
                "分母证据": denominator_evidence[patient_id]["citation"],
                "病种": disease,
                "是否追加评估": added_evaluation,
                "追加评估证据": " | ".join(evaluation_evidence[:4]),
                "追加评估信息增益_预标注": information_gain,
                "手术决策变化_预标注": decision_change,
                "决策证据": decision_evidence,
                "最终远处转移状态_预标注": m_state,
                "M状态证据": m_evidence,
                "开发终点标签": development_endpoint,
                "开发终点证据": development_evidence,
                "开发终点可用": "是" if development_endpoint != "结局待定" else "否",
                "无追加评估代理标签": proxy,
                "内容链置信度": content_confidence,
                "正式置信度": final_confidence,
                "未能判定字段": "报告签发/审核时间" + ("；最终M状态" if m_state == "无法判断" else ""),
            }
        )

    candidate_df = pd.DataFrame(candidates)
    # High-information provisional pilot.  Selection is deliberately enriched,
    # not a prevalence sample.
    priority = {
        "取消": 0,
        "改系统治疗": 1,
        "改探查活检": 2,
        "继续根治切除": 3,
        "无法判断": 4,
    }
    candidate_df["_priority"] = candidate_df["手术决策变化_预标注"].map(priority).fillna(9)
    candidate_df["_m_priority"] = candidate_df["最终远处转移状态_预标注"].map({"M1": 0, "M0": 1, "无法判断": 2})
    ordered = candidate_df.sort_values(["_priority", "_m_priority", "首次信号日期_检查日期暂代"])
    changed = ordered[ordered["手术决策变化_预标注"].isin(["取消", "推迟", "改系统治疗", "改探查活检"])]
    continued = ordered[ordered["手术决策变化_预标注"].eq("继续根治切除")]
    pilot = pd.concat(
        [
            changed.head(12),
            continued[continued["最终远处转移状态_预标注"].eq("M1")].head(8),
            continued[continued["最终远处转移状态_预标注"].eq("M0")].head(5),
            continued[continued["最终远处转移状态_预标注"].eq("无法判断")].head(5),
        ],
        ignore_index=True,
    ).drop_duplicates(subset=["患者ID"], keep="first")
    # If one stratum is sparse, fill only to the 30-case startup target from
    # the remaining pool.  This is an explicitly enriched development sample.
    if len(pilot) < 30:
        fill = ordered[~ordered["患者ID"].isin(pilot["患者ID"])].head(30 - len(pilot))
        pilot = pd.concat([pilot, fill], ignore_index=True)
    pilot = pilot.drop(columns=["_priority", "_m_priority"]).copy()
    candidate_df = candidate_df.drop(columns=["_priority", "_m_priority"])

    # Human review sample: all rare labels first, then a fixed sample of the
    # common labels.  It remains provisional until T0 is repaired.
    review_sample = (
        pilot.assign(
            _review_priority=pilot["手术决策变化_预标注"].map(priority).fillna(9),
            _review_m=pilot["最终远处转移状态_预标注"].map({"M1": 0, "M0": 1, "无法判断": 2}),
        )
        .sort_values(["_review_priority", "_review_m", "患者ID"])
        .drop(columns=["_review_priority", "_review_m"])
        .head(25)
        .copy()
    )
    review_sample["人工结论"] = ""
    review_sample["明显错标"] = ""
    review_sample["引用真实"] = ""
    review_sample["备注"] = ""

    weak_df = pd.DataFrame(weak_rows)
    discard_df = pd.DataFrame(
        [{"丢弃原因": reason, "数量": count} for reason, count in discard.most_common()]
    )
    blocker_df = pd.DataFrame(
        [
            {
                "被卡字段": "首次信号T0",
                "状态": "阻断",
                "原因": "原始影像251个CSV分片仅有检查日期/检查时间，无签发或审核时间；工作区未找到可回填源。",
                "所需数据": "RIS报告签发时间或审核时间，至少包含患者ID/影像号/检查号及签发时间。",
            },
            {
                "被卡字段": "最终M状态",
                "状态": "部分缺失",
                "原因": "严格M0仅接受同部位阴性病理或影像明确排除；普通≥6个月影像随访只进入代理阴性层，不冒充M0。",
                "所需数据": "远处病灶病理、影像明确排除，或用于开发层的首次信号后≥6个月可评估影像。",
            },
        ]
    )

    candidate_df.to_parquet(OUT / "strong_pool_pre_t0.parquet", index=False)
    weak_df.to_parquet(OUT / "weak_pool.parquet", index=False)
    pilot.to_parquet(OUT / "pilot_candidates_pre_t0.parquet", index=False)

    label_distribution = Counter(pilot["手术决策变化_预标注"].astype(str))
    dominant_share = max(label_distribution.values(), default=0) / max(len(pilot), 1)
    high_confidence_share = (pilot["正式置信度"] == "高").mean() if len(pilot) else 0.0
    m_determinable_share = (pilot["最终远处转移状态_预标注"] != "无法判断").mean() if len(pilot) else 0.0
    development_endpoint_share = (pilot["开发终点可用"] == "是").mean() if len(pilot) else 0.0
    manifest = {
        "status": "BLOCKED_T0",
        "generated_at": "2026-09-26",
        "scope": "高纯度决策窗口富集开发集预筛；非发生率队列",
        "source_report_count": int(len(image)),
        "strong_report_hits_after_text_filters": int(len(report_signal)),
        "strong_patients_with_prior_independent_pancreas_path": int(len(patient_signals)),
        "strong_pool_pre_t0_n": int(len(candidate_df)),
        "m0_recovery_candidate_n": int((candidate_df["最终远处转移状态_预标注"] == "M0").sum()),
        "added_evaluation_recovery_candidate_n": int((candidate_df["是否追加评估"] == "有").sum()),
        "pilot_pre_t0_n": int(len(pilot)),
        "weak_pool_n": int(len(weak_df)),
        "discard_reason_counts": dict(discard),
        "decision_label_counts_in_pilot": dict(label_distribution),
        "dominant_decision_label_share": float(dominant_share),
        "break_chain": bool(dominant_share > 0.70),
        "gates": {
            "gate1_high_confidence_share": float(high_confidence_share),
            "gate1_pass": False,
            "gate1_note": "正式置信度因T0缺失统一降为低，禁止把内容链置信度替代正式置信度。",
            "gate2_independent_rerun_consistency": None,
            "gate2_pass": False,
            "gate2_note": "T0修复前不执行独立重跑。",
            "gate3_m_status_determinable_share": float(m_determinable_share),
            "gate3_strict_pass": bool(m_determinable_share >= 0.50),
            "gate3_development_endpoint_observable_share": float(development_endpoint_share),
            "gate3_pass": bool(development_endpoint_share >= 0.50),
            "gate3_note": "开发闸门允许≥6月可评估影像作为代理阴性，但最终M状态仍保持无法判断，不计作M0。",
        },
        "human_review": {
            "planned_n": int(len(review_sample)),
            "completed_n": 0,
            "wrong_label_n": None,
            "missing_or_false_citation_n": None,
            "note": "已生成抽检表；正式T0修复前不宣告人工抽检通过。",
        },
        "outputs": [
            "pilot_screening_pre_t0.xlsx",
            "strong_pool_pre_t0.parquet",
            "pilot_candidates_pre_t0.parquet",
            "weak_pool.parquet",
        ],
    }
    (OUT / "screening_manifest_pre_t0.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    with pd.ExcelWriter(OUT / "pilot_screening_pre_t0.xlsx", engine="openpyxl") as writer:
        pilot.to_excel(writer, index=False, sheet_name="预备集_非定稿")
        review_sample.to_excel(writer, index=False, sheet_name="人工抽检25例")
        candidate_df[candidate_df["最终远处转移状态_预标注"].eq("M0")].to_excel(
            writer, index=False, sheet_name="M0待补决策文书"
        )
        candidate_df[candidate_df["是否追加评估"].eq("有")].to_excel(
            writer, index=False, sheet_name="追加评估待补决策"
        )
        discard_df.to_excel(writer, index=False, sheet_name="丢弃原因")
        blocker_df.to_excel(writer, index=False, sheet_name="阻断字段")
        pd.DataFrame(
            [{"指标": key, "值": json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value} for key, value in manifest.items()]
        ).to_excel(writer, index=False, sheet_name="预筛清单")
        style_workbook(writer)

    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
