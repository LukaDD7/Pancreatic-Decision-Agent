import csv
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from scripts.cohort_construction.cohort_100.step02_build_closed_decision_chains import (
    DISTANT_TREATMENT,
    PAIR,
    RADICAL,
    SURGERY_DATE,
    classify_actual,
    compact,
    extract_intraoperative_finding,
    parse_json_list,
)
from scripts.cohort_construction.cohort_100.step01_build_pre_t0_signal_pool import (
    POST_TREATMENT,
    extract_signal,
    report_citation,
    site_name,
)
from scripts.cohort_construction.paths import data_root


ROOT = data_root()
PROJECT = ROOT / "pipeline_outputs_stage8_v1" / "pancreas_decision_window_enriched_v1"
OUT = PROJECT / "restricted"
INDEX = (
    ROOT
    / "pipeline_outputs_stage8_v1"
    / "pancreas_imaging_report_index_v2"
    / "restricted"
    / "index"
    / "report_index.parquet"
)
DOCUMENTS = ROOT / "患者入院出院文书.csv"
STRONG_POOL = OUT / "strong_pool_pre_t0.parquet"
OLD_60 = OUT / "closed_decision_chain_cohort_60_pre_t0.parquet"
SHORTLIST = (
    ROOT
    / "pipeline_outputs_stage8_v1"
    / "pancreas_decision_failure_cases_v3"
    / "restricted"
    / "shortlist.xlsx"
)

PLAN_RADICAL = re.compile(
    r"(?:拟施手术|拟施手术名称和手术方式|拟实施|拟行|拟施)\s*[:：]?\s*"
    r"(?P<planned>[^。；\n]{0,160}(?:胰十二指肠切除|Whipple|胰体尾[^。；\n]{0,30}切除|"
    r"远端胰腺切除|胰腺次全切除|全胰切除|胰腺根治|胰体癌根治|胰尾癌根治)[^。；\n]{0,120})"
)
ACTUAL_NAME = re.compile(
    r"(?:实施手术名称|实施手术|手术名称)\s*[:：]\s*(?P<actual>[^。；\n]{1,240})"
)
POSTOP_ACTUAL = re.compile(
    r"患者今日在[^。；\n]{0,30}(?:麻醉下)?行(?P<actual>[^。；\n]{1,220}(?:术|切除|活检)[^。；\n]{0,80})"
)
STAGING_LAP = re.compile(
    r"分期腹腔镜|腹腔镜下?(?:腹腔)?探查|腹腔镜探查|腹腔镜[^。；\n]{0,35}(?:活检|探查)"
)
OPEN_EXPLORE = re.compile(r"剖腹探查|开腹探查")
EXPLICIT_CANCEL = re.compile(
    r"取消手术|暂缓手术|放弃手术|未行根治性切除|未予根治性切除|"
    r"终止根治性切除|术中终止手术"
)


def style_workbook(writer):
    for worksheet in writer.book.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for index, column in enumerate(worksheet.columns, 1):
            values = [str(cell.value or "") for cell in list(column)[:200]]
            width = min(max(max((len(value) for value in values), default=8) + 2, 10), 52)
            worksheet.column_dimensions[get_column_letter(index)].width = width


def build_signal_base():
    image = pd.read_parquet(INDEX)
    image["patient_id"] = image["patient_id"].astype(str)
    image["_exam_dt"] = pd.to_datetime(image["exam_datetime"], errors="coerce")
    image = image.sort_values(["patient_id", "_exam_dt"])
    existing = pd.read_parquet(STRONG_POOL).set_index("患者ID", drop=False)
    rows = {}
    for _, report in image.iterrows():
        patient_id = str(report["patient_id"])
        if patient_id in rows or pd.isna(report["_exam_dt"]):
            continue
        conclusion = str(report.get("description_deidentified") or "")
        quote, _ = extract_signal(conclusion)
        if not quote or POST_TREATMENT.search(conclusion):
            continue
        rows[patient_id] = {
            "患者ID": patient_id,
            "首次信号日期_检查日期暂代": report["_exam_dt"].normalize(),
            "信号等级": "强可疑",
            "部位": site_name(quote),
            "信号原文": quote,
            "信号证据引用": report_citation(report, quote),
            "病种": "其他/待定",
            "最终远处转移状态": "待核",
            "M状态证据": "",
            "来源": "全量影像强信号扩展筛选",
            "原预筛分母状态": "首次信号前未确认独立胰腺路径；由手术计划反向确认",
            "旧shortlist_case_id": "",
        }
    for patient_id, row in existing.iterrows():
        if patient_id not in rows:
            continue
        rows[patient_id].update(
            {
                "病种": row["病种"],
                "最终远处转移状态": row["最终远处转移状态_预标注"],
                "M状态证据": row["M状态证据"],
                "来源": "原全量强信号预筛",
                "原预筛分母状态": "首次信号前已有独立胰腺外科评估路径证据",
            }
        )

    shortlist = pd.read_excel(SHORTLIST, sheet_name="五类特殊病例")
    for _, row in shortlist.iterrows():
        patient_id = str(row["patient_id"])
        events = parse_json_list(row["preop_strong_metastasis_json"]) + parse_json_list(
            row["preop_uncertain_metastasis_json"]
        )
        event_date = pd.to_datetime(row["event_date"], errors="coerce")
        valid = []
        for event in events:
            date = pd.to_datetime(event.get("date"), errors="coerce")
            if pd.notna(date) and pd.notna(event_date) and date < event_date <= date + pd.Timedelta(days=90):
                valid.append((date, event))
        if not valid:
            continue
        date, event = sorted(valid, key=lambda item: item[0])[0]
        text = str(event.get("text") or "")
        explicit = bool(
            re.search(r"(?:肝|腹膜|网膜|肺|骨|肾上腺)[^。；\n]{0,20}(?:转移瘤|多发转移|转移灶|转移)", text)
            and not re.search(r"待排|可能|不排除|不除外|可疑", text)
        )
        citation = f"{event.get('source_record_key','')}@{date.date()}:{compact(text)}"
        if patient_id in rows:
            rows[patient_id]["旧shortlist_case_id"] = str(row.get("case_id") or "")
            if explicit:
                rows[patient_id]["信号等级"] = "明确"
            continue
        rows[patient_id] = {
            "患者ID": patient_id,
            "首次信号日期_检查日期暂代": date.normalize(),
            "信号等级": "明确" if explicit else "强可疑",
            "部位": "待人工确认",
            "信号原文": compact(text),
            "信号证据引用": citation,
            "病种": str(row.get("disease_labels") or "其他/待定"),
            "最终远处转移状态": "待核",
            "M状态证据": "",
            "来源": "旧shortlist补充",
            "原预筛分母状态": "旧shortlist已有胰腺外科候选路径",
            "旧shortlist_case_id": str(row.get("case_id") or ""),
        }
    return rows


def actual_from_text(text):
    match = ACTUAL_NAME.search(text)
    if match:
        return compact(match.group("actual"), 320), match
    match = POSTOP_ACTUAL.search(text)
    if match:
        return compact(match.group("actual"), 320), match
    return "", None


def main():
    signals = build_signal_base()
    patient_ids = set(signals)
    paired = defaultdict(list)
    plans = defaultdict(list)
    actuals = defaultdict(list)
    cancellation_evidence = defaultdict(list)
    csv.field_size_limit(2**31 - 1)
    with DOCUMENTS.open("r", encoding="gb18030", errors="replace", newline="") as handle:
        reader = csv.DictReader(handle)
        for row_number, row in enumerate(reader, start=2):
            patient_id = str(row.get("PATIENT_ID") or "")
            if patient_id not in patient_ids:
                continue
            create_time = pd.to_datetime(row.get("CREATE_DATE_TIME"), errors="coerce")
            if pd.isna(create_time):
                continue
            signal_date = signals[patient_id]["首次信号日期_检查日期暂代"]
            days = (create_time.normalize() - signal_date.normalize()).days
            if not (0 <= days <= 100):
                continue
            text = str(row.get("文书内容") or "")
            citation_base = f"患者入院出院文书.csv:{row_number}@{create_time.date()}"
            pair_match = PAIR.search(text)
            if pair_match:
                planned = compact(pair_match.group("planned"), 240)
                actual = compact(pair_match.group("actual"), 320)
                date_match = SURGERY_DATE.search(text)
                operation_date = (
                    pd.Timestamp(int(date_match.group(1)), int(date_match.group(2)), int(date_match.group(3)))
                    if date_match
                    else create_time.normalize()
                )
                operation_days = (operation_date.normalize() - signal_date.normalize()).days
                if 1 <= operation_days <= 90:
                    intraop_category, intraop_text = extract_intraoperative_finding(text)
                    paired[patient_id].append(
                        {
                            "date": operation_date,
                            "days": operation_days,
                            "planned": planned,
                            "actual": actual,
                            "plan_citation": f"{citation_base}:拟实施手术名称：{planned}",
                            "actual_citation": f"{citation_base}:实施手术名称：{actual}",
                            "intraop_category": intraop_category,
                            "intraop_citation": f"{citation_base}:{intraop_text}" if intraop_text else "",
                            "explicit_cancel": bool(EXPLICIT_CANCEL.search(text)),
                        }
                    )
            plan_match = PLAN_RADICAL.search(text)
            if plan_match:
                plans[patient_id].append(
                    {
                        "date": create_time.normalize(),
                        "days": days,
                        "planned": compact(plan_match.group("planned"), 240),
                        "citation": f"{citation_base}:{compact(text[max(0, plan_match.start()-100):plan_match.end()+260])}",
                    }
                )
            actual, actual_match = actual_from_text(text)
            if actual and (RADICAL.search(actual) or STAGING_LAP.search(actual) or OPEN_EXPLORE.search(actual) or re.search(r"活检", actual)):
                date_match = SURGERY_DATE.search(text)
                operation_date = (
                    pd.Timestamp(int(date_match.group(1)), int(date_match.group(2)), int(date_match.group(3)))
                    if date_match
                    else create_time.normalize()
                )
                operation_days = (operation_date.normalize() - signal_date.normalize()).days
                if 1 <= operation_days <= 90:
                    intraop_category, intraop_text = extract_intraoperative_finding(text)
                    actuals[patient_id].append(
                        {
                            "date": operation_date,
                            "days": operation_days,
                            "actual": actual,
                            "citation": f"{citation_base}:{compact(text[max(0, actual_match.start()-80):actual_match.end()+260])}",
                            "intraop_category": intraop_category,
                            "intraop_citation": f"{citation_base}:{intraop_text}" if intraop_text else "",
                            "staging_laparoscopy": bool(STAGING_LAP.search(actual)),
                        }
                    )
            cancel_match = EXPLICIT_CANCEL.search(text)
            document_name = str(row.get("文书名称") or "")
            if cancel_match and not re.search(r"知情同意|术前讨论|术前小结", document_name):
                cancellation_evidence[patient_id].append(
                    f"{citation_base}:{compact(text[max(0, cancel_match.start()-120):cancel_match.end()+260])}"
                )

    exclusions = Counter()
    rows = []
    for patient_id, signal in signals.items():
        operation = None
        if paired.get(patient_id):
            operation = sorted(paired[patient_id], key=lambda item: (item["days"], item["date"]))[0]
        else:
            patient_plans = [item for item in plans.get(patient_id, []) if RADICAL.search(item["planned"])]
            patient_actuals = actuals.get(patient_id, [])
            compatible = []
            for actual in patient_actuals:
                prior_plans = [plan for plan in patient_plans if plan["date"] <= actual["date"]]
                if prior_plans:
                    plan = sorted(prior_plans, key=lambda item: item["date"])[-1]
                    compatible.append((actual, plan))
            if compatible:
                actual, plan = sorted(compatible, key=lambda pair: (pair[0]["days"], pair[0]["date"]))[0]
                operation = {
                    **actual,
                    "planned": plan["planned"],
                    "plan_citation": plan["citation"],
                    "actual_citation": actual["citation"],
                    "explicit_cancel": bool(cancellation_evidence.get(patient_id)),
                }
        if not operation:
            exclusions["缺少可配对的根治计划与实际处置"] += 1
            continue
        if not RADICAL.search(operation["planned"]):
            exclusions["计划术式不是明确胰腺根治性切除"] += 1
            continue
        actual_class = classify_actual(operation["actual"])
        if not operation.get("intraop_citation") and re.search(
            r"转移|腹腔继发恶性肿瘤", operation.get("actual_citation", "")
        ):
            operation["intraop_category"] = "术中发现远处转移/高度可疑远处病灶"
            operation["intraop_citation"] = operation["actual_citation"]
        staging_lap = bool(STAGING_LAP.search(operation["actual"])) and not RADICAL.search(operation["actual"])
        if staging_lap:
            pathway = "计划根治→分期腹腔镜→取消根治切除"
        elif actual_class == "根治性切除并同期处理远处病灶":
            pathway = "疑似转移→根治切除并同期处理远处病灶"
        elif actual_class == "根治性切除":
            pathway = "疑似转移→继续单纯根治切除"
        elif actual_class in {"探查/活检", "探查/活检并姑息处理", "姑息处理"}:
            pathway = "疑似转移→手术继续但降级"
        else:
            exclusions["实际处置无法归入四类"] += 1
            continue
        rows.append(
            {
                **signal,
                "计划术式": operation["planned"],
                "计划证据引用": operation["plan_citation"],
                "实际处置日期": operation["date"].date(),
                "距首次信号天数": operation["days"],
                "实际术式/处置": operation["actual"],
                "切除性质": actual_class,
                "实际处置证据引用": operation["actual_citation"],
                "术中发现分类": operation["intraop_category"],
                "术中发现证据引用": operation["intraop_citation"],
                "明确取消/终止语义": "有" if operation.get("explicit_cancel") or cancellation_evidence.get(patient_id) else "无",
                "取消/终止证据引用": " | ".join(cancellation_evidence.get(patient_id, [])[:3]),
                "四类路径": pathway,
                "根治切除实施状态": "已实施" if actual_class in {"根治性切除", "根治性切除并同期处理远处病灶"} else "未实施",
                "决策链状态": "闭合",
                "T0状态": "缺报告签发/审核时间；当前以检查时间暂代",
            }
        )

    all_closed = pd.DataFrame(rows).drop_duplicates("患者ID")
    # The user requested expansion on top of the reviewed 60-case cohort.
    # Preserve every prior case even when the broader first-signal selection
    # chooses an earlier report and therefore misses the old 90-day window.
    old = pd.read_parquet(OLD_60).copy()
    old_path = {
        "继续根治切除": "疑似转移→继续单纯根治切除",
        "根治性切除并同期处理远处病灶": "疑似转移→根治切除并同期处理远处病灶",
        "正确降级为探查/活检/姑息": "疑似转移→手术继续但降级",
    }
    old_rows = pd.DataFrame(
        {
            "患者ID": old["患者ID"],
            "首次信号日期_检查日期暂代": old["首次信号日期_检查日期暂代"],
            "信号等级": old["信号等级"],
            "部位": old["部位"],
            "信号原文": old["信号原文"],
            "信号证据引用": old["信号证据引用"],
            "病种": old["病种"],
            "最终远处转移状态": old["最终远处转移状态"],
            "M状态证据": old["M状态证据"],
            "来源": old["来源"].astype(str) + "+原60例保留",
            "原预筛分母状态": "首次信号前已有独立胰腺外科评估路径证据",
            "旧shortlist_case_id": old["旧shortlist_case_id"],
            "计划术式": old["计划术式"],
            "计划证据引用": old["计划证据引用"],
            "实际处置日期": old["实际手术日期"],
            "距首次信号天数": old["距首次信号天数"],
            "实际术式/处置": old["实际术式"],
            "切除性质": old["切除性质"],
            "实际处置证据引用": old["实际术式证据引用"],
            "术中发现分类": old["术中发现分类"],
            "术中发现证据引用": old["术中发现证据引用"],
            "明确取消/终止语义": "无",
            "取消/终止证据引用": "",
            "四类路径": old["手术决策结果"].map(old_path),
            "根治切除实施状态": old["切除性质"].map(
                lambda value: "已实施" if value in {"根治性切除", "根治性切除并同期处理远处病灶"} else "未实施"
            ),
            "决策链状态": "闭合",
            "T0状态": old["T0状态"],
        }
    )
    all_closed = pd.concat(
        [all_closed, old_rows[~old_rows["患者ID"].isin(all_closed["患者ID"])]],
        ignore_index=True,
    ).drop_duplicates("患者ID")
    order = {
        "计划根治→分期腹腔镜→取消根治切除": 0,
        "疑似转移→手术继续但降级": 1,
        "疑似转移→根治切除并同期处理远处病灶": 2,
        "疑似转移→继续单纯根治切除": 3,
    }
    all_closed["_order"] = all_closed["四类路径"].map(order)
    all_closed = all_closed.sort_values(["_order", "信号等级", "首次信号日期_检查日期暂代", "患者ID"])
    old_ids = set(old["患者ID"].astype(str))
    selected = all_closed[all_closed["患者ID"].isin(old_ids)].copy()
    staging_new = all_closed[
        all_closed["四类路径"].eq("计划根治→分期腹腔镜→取消根治切除")
        & ~all_closed["患者ID"].isin(selected["患者ID"])
    ]
    selected = pd.concat([selected, staging_new], ignore_index=True).drop_duplicates("患者ID")
    # Fill the remaining slots round-robin from the three operative pathways,
    # always adding to the currently smallest pathway.
    remaining = all_closed[~all_closed["患者ID"].isin(selected["患者ID"])].copy()
    operative_labels = [
        "疑似转移→手术继续但降级",
        "疑似转移→根治切除并同期处理远处病灶",
        "疑似转移→继续单纯根治切除",
    ]
    queues = {
        label: list(remaining[remaining["四类路径"].eq(label)].index)
        for label in operative_labels
    }
    picked = []
    while len(selected) + len(picked) < 100:
        counts = pd.concat([selected, all_closed.loc[picked] if picked else all_closed.iloc[0:0]])[
            "四类路径"
        ].value_counts()
        available = [label for label in operative_labels if queues[label]]
        if not available:
            break
        label = min(available, key=lambda value: (counts.get(value, 0), operative_labels.index(value)))
        picked.append(queues[label].pop(0))
    if picked:
        selected = pd.concat([selected, all_closed.loc[picked]], ignore_index=True)
    selected = selected.drop(columns=["_order"], errors="ignore")
    all_closed = all_closed.drop(columns=["_order"], errors="ignore")

    review_parts = []
    for label, n in (
        ("计划根治→分期腹腔镜→取消根治切除", 6),
        ("疑似转移→手术继续但降级", 8),
        ("疑似转移→根治切除并同期处理远处病灶", 8),
        ("疑似转移→继续单纯根治切除", 8),
    ):
        review_parts.append(selected[selected["四类路径"].eq(label)].head(n))
    review = pd.concat(review_parts, ignore_index=True).head(30).copy()
    if len(review) < 30:
        review_fill = selected[~selected["患者ID"].isin(review["患者ID"])].head(30 - len(review))
        review = pd.concat([review, review_fill], ignore_index=True)
    for column in ("信号引用真实", "计划术式正确", "实际处置正确", "路径标签正确", "人工结论", "备注"):
        review[column] = ""
    distribution = selected["四类路径"].value_counts().to_dict()
    dominant_share = max(distribution.values(), default=0) / max(len(selected), 1)
    manifest = {
        "status": "PRE_T0_HUMAN_REVIEW_REQUIRED",
        "generated_at": "2026-10-03",
        "all_strong_signal_patients": len(signals),
        "all_closed_four_pathway_n": len(all_closed),
        "selected_n": len(selected),
        "pathway_distribution": distribution,
        "dominant_pathway_share": dominant_share,
        "break_chain": dominant_share > 0.70,
        "exclusion_counts": dict(exclusions),
        "selection_note": "富集开发集，不用于发生率估计；弱信号未纳入。",
        "t0_blocker": "缺少影像报告签发/审核时间，首次信号暂以检查时间排序。",
    }

    selected.to_parquet(OUT / "four_pathway_cohort_100_pre_t0.parquet", index=False)
    all_closed.to_parquet(OUT / "all_closed_four_pathway_cases_pre_t0.parquet", index=False)
    (OUT / "four_pathway_cohort_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    with pd.ExcelWriter(OUT / "four_pathway_cohort_100_pre_t0.xlsx", engine="openpyxl") as writer:
        selected.to_excel(writer, index=False, sheet_name="四类集合")
        for label, sheet in (
            ("计划根治→分期腹腔镜→取消根治切除", "分期腹腔镜取消"),
            ("疑似转移→继续单纯根治切除", "继续根治"),
            ("疑似转移→手术继续但降级", "手术降级"),
            ("疑似转移→根治切除并同期处理远处病灶", "同期远端处理"),
        ):
            selected[selected["四类路径"].eq(label)].to_excel(writer, index=False, sheet_name=sheet)
        review.to_excel(writer, index=False, sheet_name="人工抽检30例")
        pd.DataFrame([{"排除原因": key, "数量": value} for key, value in exclusions.most_common()]).to_excel(
            writer, index=False, sheet_name="排除原因"
        )
        pd.DataFrame(
            [
                {"路径": "计划根治→分期腹腔镜→取消根治切除", "定义": "有明确根治计划，实际仅行腹腔镜探查/活检，未完成胰腺根治切除"},
                {"路径": "疑似转移→继续单纯根治切除", "定义": "实际完成胰腺根治切除，未见同期远处病灶切除/消融"},
                {"路径": "疑似转移→手术继续但降级", "定义": "实际改为开腹探查、活检或姑息处理"},
                {"路径": "疑似转移→根治切除并同期处理远处病灶", "定义": "完成胰腺根治切除，同时切除、剜除或消融远处病灶"},
            ]
        ).to_excel(writer, index=False, sheet_name="分类定义")
        pd.DataFrame(
            [{"指标": key, "值": json.dumps(value, ensure_ascii=False) if isinstance(value, dict) else value} for key, value in manifest.items()]
        ).to_excel(writer, index=False, sheet_name="运行清单")
        style_workbook(writer)
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
